import concurrent.futures
import base64
import gzip
import hashlib
import http.client
import io
import json
import sys
import tempfile
import unittest
import urllib.error
import urllib.request
import zipfile
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import requests
from huggingface_hub import CommitOperationAdd, CommitOperationDelete
from huggingface_hub.errors import HfHubHTTPError, RepositoryNotFoundError

from scripts import build_reader_assets_index, convert_reader_assets, pdf_assets, publish_reader_assets
from scripts import gc_reader_bucket, prune_reader_assets, publish_search_reader_index
from scripts import epub_chapters, reader_assets, scan_reader_assets


class ReaderAssetContractTests(unittest.TestCase):
    def test_bucket_gc_zero_limit_means_unlimited(self):
        self.assertEqual(gc_reader_bucket.apply_limit(["b", "a", "c"], 0), ["a", "b", "c"])
        self.assertEqual(gc_reader_bucket.apply_limit(["b", "a", "c"], 2), ["a", "b"])

    def test_bucket_gc_can_use_existing_compact_sidecar_without_new_indexes(self):
        sidecar = gzip.compress(json.dumps({"v": 1, "f": {
            "live": {"s": 2, "p": "objects/live", "b": "vomebook/pdf-pages"},
            "optimized": {"s": 2, "p": "objects/optimized", "b": "vomebook/pdf-optimized"},
        }}).encode())
        store = Mock()
        store.read_bytes.return_value = sidecar
        files = {"reader-index/reader_assets.json.gz", "objects/live", "objects/old"}
        payloads = gc_reader_bucket.read_index_payloads(store, files)
        references = gc_reader_bucket.current_references(files, {}, payloads)
        self.assertIn("objects/live", references)
        self.assertNotIn("objects/old", references)

    def test_bucket_gc_expands_page_manifest_to_derived_pages(self):
        store = Mock()
        store.read_bytes.side_effect = lambda _bucket, path: json.dumps({
            "kind": "pdf-pages", "page_count": 2,
        }).encode() if path.endswith("page-manifest.json") else b"{}"
        manifest = "objects/aa/book/pageset/page-manifest.json"
        files = {
            manifest,
            "objects/aa/book/pageset/pages/page-000001.webp",
            "objects/aa/book/pageset/pages/page-000002.webp",
            "objects/aa/book/pageset/pages/page-000003.webp",
        }
        references = {manifest}
        gc_reader_bucket.expand_reference_closure(store, files, references)
        self.assertIn("objects/aa/book/pageset/pages/page-000001.webp", references)
        self.assertIn("objects/aa/book/pageset/pages/page-000002.webp", references)
        self.assertNotIn("objects/aa/book/pageset/pages/page-000003.webp", references)

    def test_bucket_gc_expands_chapter_manifest_relative_paths(self):
        store = Mock()
        store.read_bytes.return_value = json.dumps({
            "kind": "ebook-chapters", "chapters": [{"path": "chapters/chapter-0001.xhtml"}],
            "search_index": {"path": "epub-search-index.json.gz"},
        }).encode()
        manifest = "ebook-chapters/objects/ab/book/epub-chapters/chapter-manifest.json"
        files = {
            manifest,
            "ebook-chapters/objects/ab/book/epub-chapters/chapters/chapter-0001.xhtml",
            "ebook-chapters/objects/ab/book/epub-chapters/epub-search-index.json.gz",
        }
        references = {manifest}
        gc_reader_bucket.expand_reference_closure(store, files, references)
        self.assertIn("ebook-chapters/objects/ab/book/epub-chapters/chapters/chapter-0001.xhtml", references)
        self.assertIn("ebook-chapters/objects/ab/book/epub-chapters/epub-search-index.json.gz", references)

    def test_bucket_gc_tracks_explicit_staging_bucket_paths(self):
        sidecar = {"v": 1, "f": {
            "optimized": {"s": 2, "p": "objects/optimized", "b": "vomebook/pdf-optimized"},
        }}
        references = set()
        gc_reader_bucket.collect_bucket_paths(sidecar, "vomebook/pdf-optimized", references)
        self.assertEqual(references, {"objects/optimized"})

    def test_bucket_gc_jxl_only_filters_candidates_without_touching_other_assets(self):
        store = Mock()
        store.list_files.side_effect = [
            {"reader-index/reader_assets.json.gz", "objects/a/page.webp", "objects/b/old.jxl"},
            set(),
        ]
        sidecar = gzip.compress(json.dumps({"v": 1, "f": {}}).encode())
        store.read_bytes.return_value = sidecar
        updated, expired, counts = gc_reader_bucket.plan_gc(store, 0, 0, jxl_only=True)
        self.assertEqual(counts[reader_assets.READER_ASSETS_BUCKET], 1)
        self.assertEqual(expired[reader_assets.READER_ASSETS_BUCKET], ["objects/b/old.jxl"])
        self.assertEqual(counts["__jxl_total"], 1)
        self.assertEqual(counts["__jxl_referenced"], 0)

    def test_bucket_gc_force_jxl_delete_selects_only_reader_jxl(self):
        store = Mock()
        store.list_files.side_effect = [
            {"reader-index/reader_assets.json.gz", "objects/a/live.webp", "objects/b/live.jxl"},
            {"objects/c/staging.jxl"},
        ]
        store.read_bytes.return_value = gzip.compress(json.dumps({"v": 1, "f": {}}).encode())
        _updated, expired, _counts = gc_reader_bucket.plan_gc(
            store, 0, 0, force_jxl_delete=True
        )
        self.assertEqual(expired[reader_assets.READER_ASSETS_BUCKET], ["objects/b/live.jxl"])
        self.assertEqual(expired[reader_assets.READER_STAGING_BUCKET], [])

    def test_bucket_gc_can_use_separate_input_account_credentials(self):
        class FakeBoto3:
            def __init__(self):
                self.calls = []

            def client(self, *args, **kwargs):
                self.calls.append((args, kwargs))
                return object()

        store = gc_reader_bucket.S3BucketStore.__new__(gc_reader_bucket.S3BucketStore)
        store._boto3 = FakeBoto3()
        store._config = object()
        store._access_key = "main-key"
        store._secret_key = "main-secret"
        store._namespace = "vomebook"
        store._input_namespace = "other-account"
        store._input_bucket = "melsm"
        store._clients = {}
        with patch.dict("os.environ", {
            "HF_S3_INPUT_ACCESS_KEY_ID": "",
            "HF_S3_INPUT_SECRET_ACCESS_KEY": "",
        }):
            with self.assertRaises(RuntimeError):
                store._client("other-account")

        with patch.dict("os.environ", {
            "HF_S3_INPUT_ACCESS_KEY_ID": "input-key",
            "HF_S3_INPUT_SECRET_ACCESS_KEY": "input-secret",
        }):
            store._clients = {}
            store._client("other-account")
        kwargs = store._boto3.calls[-1][1]
        self.assertEqual(kwargs["aws_access_key_id"], "input-key")

    def test_conversion_set_excludes_unsafe_or_native_media(self):
        self.assertEqual(
            set(reader_assets.CONVERTIBLE_EXTENSIONS),
            {"doc", "docx", "epub", "htm", "html", "mobi", "azw3", "fb2", "odt", "rtf", "chm", "tif", "tiff", "djvu",
             "ppt", "pptx", "pps", "odp", "xls", "xlsx", "csv", "ods", "wps", "mht", "mhtml", "ps", "caj", "kdh",
             "ape", "wma", "amr", "flv", "f4v", "rm", "rmvb", "mkv", "avi", "mpg",
             "mpeg", "mts", "ts", "wmv"},
        )
        for extension in ("pdg", "swf", "asx", "dat", "mp3", "mp4", "wav", "m4a", "flac", "mov", "mpga"):
            self.assertNotIn(extension, reader_assets.CONVERTIBLE_EXTENSIONS)

    def test_source_url_pins_revision_and_encodes_path(self):
        self.assertEqual(
            reader_assets.source_url("VoiceOfML/Test", "abc123", "目录/a b.docx"),
            "https://huggingface.co/datasets/VoiceOfML/Test/resolve/abc123/"
            "%E7%9B%AE%E5%BD%95/a%20b.docx",
        )

    def test_every_supported_extension_has_an_explicit_contract(self):
        expected = {
            "doc": ("libreoffice-docx-v2", "docx", "document.docx"),
            "docx": ("docx-native-v2", "docx", "document.docx"),
            "htm": ("sanitized-html-v5", "html", "document.html"),
            "html": ("sanitized-html-v5", "html", "document.html"),
            "epub": ("foliate-original-v1", "foliate", "document.epub"),
            "mobi": ("foliate-original-v1", "foliate", "document.mobi"),
            "azw3": ("foliate-original-v1", "foliate", "document.azw3"),
            "fb2": ("foliate-original-v1", "foliate", "document.fb2"),
            "odt": ("calibre-odt-html-v1", "html", "document.html"),
            "rtf": ("calibre-rtf-html-v1", "html", "document.html"),
            "chm": ("calibre-chm-epub-v3", "epub", "document.epub"),
            "tif": ("pillow-pdf-v2", "pdf", "document.pdf"),
            "tiff": ("pillow-pdf-v2", "pdf", "document.pdf"),
            "djvu": ("djvulibre-pdf-v2", "pdf", "document.pdf"),
            "ppt": ("libreoffice-pdf-office-v2", "pdf", "document.pdf"),
            "pptx": ("libreoffice-pdf-office-v2", "pdf", "document.pdf"),
            "pps": ("libreoffice-pdf-office-v2", "pdf", "document.pdf"),
            "odp": ("libreoffice-pdf-office-v2", "pdf", "document.pdf"),
            "xls": ("libreoffice-pdf-office-v2", "pdf", "document.pdf"),
            "xlsx": ("libreoffice-pdf-office-xlsx-v3", "pdf", "document.pdf"),
            "csv": ("libreoffice-pdf-office-v2", "pdf", "document.pdf"),
            "ods": ("libreoffice-pdf-office-v2", "pdf", "document.pdf"),
            "wps": ("libreoffice-pdf-office-v2", "pdf", "document.pdf"),
            "mht": ("sanitized-mhtml-v6", "html", "document.html"),
            "mhtml": ("sanitized-mhtml-v6", "html", "document.html"),
            "ps": ("ghostscript-pdf-v5", "pdf", "document.pdf"),
            "caj": ("caj-family-pdf-v1", "pdf", "document.pdf"),
            "kdh": ("caj-family-pdf-v1", "pdf", "document.pdf"),
            "ape": ("ffmpeg-audio-mp3-v1", "audio", "audio.mp3"),
            "wma": ("ffmpeg-audio-mp3-v1", "audio", "audio.mp3"),
            "amr": ("ffmpeg-audio-mp3-v1", "audio", "audio.mp3"),
            **{extension: ("ffmpeg-video-mp4-h264-aac-v1", "video", "video.mp4") for extension in (
                "flv", "f4v", "rm", "rmvb", "mkv", "avi", "mpg", "mpeg", "mts", "ts", "wmv",
            )},
        }
        self.assertEqual({ext: reader_assets.conversion_contract(ext) for ext in expected}, expected)
        self.assertEqual(
            reader_assets.bucket_conversion_contract("repo", "table.xlsx", "xlsx"),
            (reader_assets.SPREADSHEET_PAGE_PROFILE, "pdf", "page-manifest.json"),
        )

    def test_large_docx_keeps_native_contract(self):
        self.assertEqual(
            reader_assets.source_conversion_contract("repo", "large.docx", "docx", 128 * 1024 * 1024),
            reader_assets.conversion_contract("docx"),
        )

    def test_large_epub_keeps_native_contract(self):
        self.assertEqual(
            reader_assets.source_conversion_contract("repo", "large.epub", "epub", 128 * 1024 * 1024),
            reader_assets.conversion_contract("epub"),
        )
        self.assertEqual(
            reader_assets.source_conversion_contract("repo", "small.epub", "epub", 1),
            reader_assets.conversion_contract("epub"),
        )

    def test_pdf_fallback_can_carry_epub_chapter_manifest(self):
        manifest = reader_assets.empty_manifest()
        manifest["files"]["repo\\0book.epub"] = {
            "status": "ready", "path": "objects/a/document.pdf",
            "chapter_manifest": "ebook-chapters/objects/a/chapter-manifest.json",
            "fallback_path": "objects/a/document.pdf", "reader_mode": "pdf",
        }
        self.assertIs(reader_assets.validate_manifest(manifest), manifest)

    def test_chapter_manifest_legacy_prefix_still_rejects_unsafe_paths(self):
        for path in (
                "ebook-chapters/objects/../outside/chapter-manifest.json",
                "ebook-chapters/objects\\outside/chapter-manifest.json",
                "ebook-chapters//objects/a/chapter-manifest.json"):
            with self.subTest(path=path), self.assertRaisesRegex(ValueError, "object path"):
                reader_assets.validate_manifest({
                    "version": 1,
                    "files": {"key": {"status": "ready", "path": "objects/a/document.epub",
                                       "reader_mode": "foliate", "chapter_manifest": path}},
                 })

    def test_legacy_bucket_chapter_manifest_remains_loadable_for_repair(self):
        base = {"status": "ready", "path": "objects/aa/document.epub", "reader_mode": "foliate"}
        manifest = reader_assets.empty_manifest()
        manifest["files"]["key"] = {
            **base, "chapter_manifest": "objects/aa/epub-chapters/chapter-manifest.json",
            "chapter_bucket": reader_assets.READER_ASSETS_BUCKET,
        }
        self.assertIs(reader_assets.validate_manifest(manifest), manifest)
        manifest["files"]["key"]["chapter_manifest"] = (
            "ebook-chapters/objects/aa/epub-chapters/chapter-manifest.json"
        )
        self.assertIs(reader_assets.validate_manifest(manifest), manifest)

    def test_passwords_require_an_explicit_marker_or_known_source(self):
        self.assertEqual(reader_assets.source_password("repo", "资料〔密码：123〕.docx"), "123")
        self.assertEqual(reader_assets.source_password("repo", "毛的遗产 密码0000.pdf"), "0000")
        self.assertEqual(reader_assets.source_password("repo", "对中帝论的批判（密码：1921）.pdf"), "1921")
        self.assertEqual(reader_assets.source_password(
            "VoiceOfML/MLMRL-Library", "基础入门书单/入门答疑/风正集.pdf",
        ), "230505")
        self.assertEqual(reader_assets.source_password(
            "VoiceOfML/MLMRL-Hub", "000269/1870520043_3072_风正集230220.pdf",
        ), "230220")
        self.assertEqual(reader_assets.source_password(
            "VoiceOfML/MLMRL-Hub", "001346/1870520043_15344_风正集230123.pdf",
        ), "230123")
        self.assertEqual(reader_assets.source_password("repo", "密码学原理.pdf"), "")
        self.assertEqual(reader_assets.source_password("repo", "book password: s3cr3t-1.pdf"), "s3cr3t-1")
        self.assertEqual(reader_assets.source_password("repo", "book passwd=p@ss.docx"), "p@ss")

    def test_manifest_rejects_unsafe_object_paths(self):
        for path in ("../../outside.pdf", "/tmp/outside.pdf", "objects/../outside.pdf", "objects\\outside.pdf"):
            with self.subTest(path=path), self.assertRaisesRegex(ValueError, "object path"):
                reader_assets.validate_manifest({
                    "version": 1, "files": {"key": {"status": "ready", "path": path}},
                })

    def test_manifest_rejects_invalid_ready_entry_metadata(self):
        base = {"status": "ready", "path": "objects/x", "reader_mode": "pdf", "bytes": 10,
                "sha256": "a" * 64}
        for field, value in (("reader_mode", "unknown"), ("bytes", 0), ("sha256", "bad")):
            with self.subTest(field=field), self.assertRaisesRegex(ValueError, "ready entry"):
                reader_assets.validate_manifest({"version": 1, "files": {"key": {**base, field: value}}})

    def test_manifest_rejects_invalid_file_status(self):
        with self.assertRaisesRegex(ValueError, "invalid status"):
            reader_assets.validate_manifest({"version": 1, "files": {"key": {"status": "pending"}}})

    def test_only_known_password_pdfs_have_a_conversion_contract(self):
        self.assertEqual(
            reader_assets.source_conversion_contract("repo", "书（密码1949）.pdf", "pdf"),
            reader_assets.PROTECTED_PDF_CONTRACT,
        )
        self.assertIsNone(reader_assets.source_conversion_contract("repo", "普通书.pdf", "pdf"))

    def test_chapter_manifest_validates_order_and_fallback(self):
        chapter = {"index": 1, "title": "第一章", "path": "chapters/chapter-0001.xhtml", "bytes": 3, "sha256": "a" * 64}
        manifest = {"version": 1, "kind": "epub-chapters", "fallback": "objects/aa/" + "b" * 64 + "/profile/document.pdf", "chapters": [chapter], "search_index": {"path": "epub-search-index.json.gz", "bytes": 3, "sha256": "c" * 64}}
        self.assertIs(reader_assets.validate_chapter_manifest(manifest), manifest)

    def test_chapter_bundle_follows_spine_and_sanitizes(self):
        with tempfile.TemporaryDirectory() as root:
            root, epub, output = Path(root), Path(root) / "book.epub", Path(root) / "bundle"
            with zipfile.ZipFile(epub, "w") as archive:
                archive.writestr("mimetype", "application/epub+zip")
                archive.writestr("META-INF/container.xml", '<container><rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>')
                archive.writestr("OEBPS/content.opf", '<package><manifest><item id="a" href="a.xhtml" media-type="application/xhtml+xml"/><item id="b" href="b.xhtml" media-type="application/xhtml+xml"/><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/><item id="i" href="images/x.png" media-type="image/png"/></manifest><spine><itemref idref="nav"/><itemref idref="a"/><itemref idref="b"/></spine></package>')
                archive.writestr("OEBPS/nav.xhtml", '<html xmlns:epub="http://www.idpf.org/2007/ops"><body><nav epub:type="toc"><ol><li><a href="a.xhtml">第一章：开始</a></li><li><a href="b.xhtml">第二章：继续</a></li></ol></nav></body></html>')
                archive.writestr("OEBPS/a.xhtml", '<html><body><script>bad()</script><img src="images/x.png"/><p>one</p></body></html>')
                archive.writestr("OEBPS/b.xhtml", '<html><body><img src="images/x.png"/><p>two</p></body></html>')
                archive.writestr("OEBPS/images/x.png", b"png")
            manifest = epub_chapters.build_bundle(epub, output)
            self.assertEqual([item["index"] for item in manifest["chapters"]], [1, 2])
            self.assertNotIn("script", (output / "chapters/chapter-0001.xhtml").read_text())
            self.assertEqual([item["title"] for item in manifest["chapters"]], ["第一章：开始", "第二章：继续"])
            self.assertTrue((output / "resources/shared/OEBPS/images/x.png").is_file())
            self.assertTrue((output / "epub-search-index.json.gz").is_file())

    def test_chm_chapter_bundle_keeps_nonspine_pages_and_fragments(self):
        with tempfile.TemporaryDirectory() as root:
            root, epub, output = Path(root), Path(root) / "book.epub", Path(root) / "bundle"
            with zipfile.ZipFile(epub, "w") as archive:
                archive.writestr("mimetype", "application/epub+zip")
                archive.writestr("META-INF/container.xml", '<container><rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>')
                archive.writestr("OEBPS/content.opf", '<package><manifest><item id="a" href="a.xhtml" media-type="application/xhtml+xml"/><item id="b" href="nested/b.xhtml" media-type="application/xhtml+xml"/><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/></manifest><spine><itemref idref="a"/></spine></package>')
                archive.writestr("OEBPS/nav.xhtml", '<html xmlns:epub="http://www.idpf.org/2007/ops"><body><nav epub:type="toc"><ol><li><a href="a.xhtml">一</a></li><li><a href="nested/b.xhtml#target">二</a></li></ol></nav></body></html>')
                archive.writestr("OEBPS/a.xhtml", '<html><body><p><a href="nested/b.xhtml#target">跳转</a></p></body></html>')
                archive.writestr("OEBPS/nested/b.xhtml", '<html><body><h1 id="target">目标</h1><p>正文</p></body></html>')
            manifest = epub_chapters.build_bundle(epub, output, include_all_documents=True)
            self.assertEqual(len(manifest["chapters"]), 2)
            first = (output / "chapters/chapter-0001.xhtml").read_text(encoding="utf-8")
            self.assertIn('href="chapter-0002.xhtml#target"', first)
            self.assertEqual(manifest["toc"][1]["chapter"], 2)
            self.assertEqual(manifest["toc"][1]["fragment"], "target")
            self.assertEqual(manifest["search_index"]["bytes"], (output / "epub-search-index.json.gz").stat().st_size)

    def test_chapter_bundle_skips_oversized_resource_sets(self):
        with tempfile.TemporaryDirectory() as root:
            epub = Path(root) / "book.epub"
            with zipfile.ZipFile(epub, "w") as archive:
                archive.writestr("mimetype", "application/epub+zip")
                archive.writestr("META-INF/container.xml", '<container><rootfiles><rootfile full-path="content.opf"/></rootfiles></container>')
                archive.writestr("content.opf", '<package><manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/><item id="image" href="image.png" media-type="image/png"/></manifest><spine><itemref idref="chapter"/></spine></package>')
                archive.writestr("chapter.xhtml", '<html><body><p>这是一个足够长的章节正文，用来验证超大资源集合不会生成独立章节包。</p><img src="image.png"/></body></html>')
                archive.writestr("image.png", b"image")

            with patch.object(epub_chapters, "MAX_CHAPTER_RESOURCE_BYTES", 1), self.assertRaisesRegex(
                    ValueError, "chapter resource budget"):
                epub_chapters.build_bundle(epub, Path(root) / "bundle")

    def test_all_native_ebooks_use_chapter_streams(self):
        for extension in ("epub", "mobi", "azw3", "fb2"):
            self.assertTrue(reader_assets.needs_epub_chapters(extension, "foliate", 1))
        self.assertTrue(reader_assets.needs_epub_chapters("mobi", "foliate", 10 ** 9))
        self.assertTrue(reader_assets.needs_epub_chapters("azw3", "foliate", 10 ** 9))
        self.assertFalse(reader_assets.needs_epub_chapters("epub", "pdf", 10 ** 9))

    def test_mid_sized_epub_is_split_for_on_demand_loading(self):
        for extension in ("epub", "mobi", "azw3", "fb2"):
            with self.subTest(extension=extension):
                self.assertTrue(reader_assets.needs_epub_chapters(
                    extension, "foliate", 8 * 1024 * 1024))
                self.assertTrue(reader_assets.needs_epub_chapters(
                    extension, "foliate", 8 * 1024 * 1024 - 1))
        self.assertTrue(reader_assets.needs_epub_chapters("chm", "epub", 1))

    def test_chapter_bundle_can_publish_text_without_resources(self):
        with tempfile.TemporaryDirectory() as root:
            root, epub, output = Path(root), Path(root) / "book.epub", Path(root) / "bundle"
            with zipfile.ZipFile(epub, "w") as archive:
                archive.writestr("mimetype", "application/epub+zip")
                archive.writestr("META-INF/container.xml", '<container><rootfiles><rootfile full-path="content.opf"/></rootfiles></container>')
                archive.writestr("content.opf", '<package><manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/><item id="image" href="image.png" media-type="image/png"/></manifest><spine><itemref idref="chapter"/></spine></package>')
                archive.writestr("chapter.xhtml", '<html><body><p>这是章节正文内容，用于全文检索。</p><img src="image.png"/></body></html>')
                archive.writestr("image.png", b"x" * 10)
            manifest = epub_chapters.build_bundle(epub, output, include_resources=False)
            self.assertTrue((output / "epub-search-index.json.gz").is_file())
            self.assertFalse((output / "resources/chapter-0001/image.png").exists())
            self.assertNotIn("resources", manifest)


class ScannerTests(unittest.TestCase):
    def setUp(self):
        self.records = [
            {"Repo": "VoiceOfML/Test", "File": "Book", "Extension": "docx", "Folder": ["A"], "Size": 10},
            {"Repo": "VoiceOfML/Test", "File": "Scan", "Extension": "pdg", "Folder": [], "Size": 20},
            {"Repo": "VoiceOfML/Test", "File": "Photo", "Extension": "jpg", "Folder": [], "Size": 30},
        ]
        self.revisions = {"VoiceOfML/Test": "rev1"}

    def test_large_epub_conversion_builds_chapter_bundle(self):
        item = {
            "key": "VoiceOfML/Test\0Big.epub", "extension": "epub", "repo": "VoiceOfML/Test",
            "path": "Big.epub", "source_url": "https://example.test/Big.epub",
            "source_revision": "rev1", "profile": "foliate-original-v1",
            "reader_mode": "foliate", "output_name": "document.epub",
        }
        with tempfile.TemporaryDirectory() as root:
            bundle = Path(root)

            def download(_url, target):
                with zipfile.ZipFile(target, "w") as archive:
                    archive.writestr("mimetype", "application/epub+zip")
                    archive.writestr("META-INF/container.xml", '<container><rootfiles><rootfile full-path="content.opf"/></rootfiles></container>')
                    archive.writestr("content.opf", '<package><manifest><item id="a" href="a.xhtml" media-type="application/xhtml+xml"/><item id="b" href="b.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="a"/><itemref idref="b"/></spine></package>')
                    archive.writestr("a.xhtml", "<html><body><p>chapter one</p></body></html>")
                    archive.writestr("b.xhtml", "<html><body><p>chapter two</p></body></html>")
                    return hashlib.sha256(target.read_bytes()).hexdigest(), 1

            def normalize(command, **_kwargs):
                Path(command[2]).write_bytes(Path(command[1]).read_bytes())

            with patch.object(convert_reader_assets, "download_source", side_effect=download), \
                    patch.object(convert_reader_assets, "run_checked", side_effect=normalize):
                result = convert_reader_assets.convert_item(item, bundle)
            self.assertEqual(result["status"], "ready")
            self.assertTrue(result["path"].endswith("/document.epub"))
            manifest_path = result["chapter_manifest"]
            self.assertTrue(manifest_path.endswith("/epub-chapters/chapter-manifest.json"))
            self.assertTrue((bundle / manifest_path).is_file())
            parent = Path(manifest_path).parent
            self.assertTrue((bundle / parent / "chapters" / "chapter-0001.xhtml").is_file())
            self.assertTrue((bundle / parent / "epub-search-index.json.gz").is_file())
            self.assertIn("chapter one", (bundle / parent / "chapters" / "chapter-0001.xhtml").read_text())

    def test_small_epub_conversion_builds_chapter_bundle(self):
        item = {
            "key": "VoiceOfML/Test\0Small.epub", "extension": "epub", "repo": "VoiceOfML/Test",
            "path": "Small.epub", "source_url": "https://example.test/Small.epub",
            "source_revision": "rev1", "profile": "foliate-original-v1",
            "reader_mode": "foliate", "output_name": "document.epub",
        }
        with tempfile.TemporaryDirectory() as root:
            def download(_url, target):
                target.write_bytes(b"small epub")
                return hashlib.sha256(b"small epub").hexdigest(), 10
            with patch.object(convert_reader_assets, "download_source", side_effect=download):
                result = convert_reader_assets.convert_item(item, Path(root))
            self.assertEqual(result["status"], "ready")
            self.assertEqual(result["chapter_bundle_profile"], reader_assets.EPUB_CHAPTER_PROFILE)
            self.assertIn("chapter_bundle_error", result)

    def test_queues_only_supported_changed_files(self):
        queue = scan_reader_assets.build_queue(
            self.records, self.revisions, reader_assets.empty_manifest()
        )
        self.assertEqual(len(queue), 1)
        self.assertEqual(queue[0]["path"], "A/Book.docx")
        self.assertEqual(queue[0]["profile"], "docx-native-v2")
        self.assertEqual(queue[0]["reader_mode"], "docx")

    def test_bucket_migration_queues_native_text_markdown_and_images(self):
        records = [
            {"Repo": "VoiceOfML/Test", "File": "Notes", "Extension": "md", "Folder": [], "Size": 10},
            {"Repo": "VoiceOfML/Test", "File": "Photo", "Extension": "jpg", "Folder": [], "Size": 10},
            {"Repo": "VoiceOfML/Test", "File": "Plain", "Extension": "txt", "Folder": [], "Size": 10},
            {"Repo": "VoiceOfML/Test", "File": "Saved Page", "Extension": "mht", "Folder": [], "Size": 10},
        ]
        queue = scan_reader_assets.build_queue(
            records, self.revisions, reader_assets.empty_manifest(), bucket_migrate=True,
        )
        self.assertEqual(
            [(item["extension"], item["reader_mode"], item["output_name"]) for item in queue],
             [("md", "markdown", "document.md"), ("jpg", "image", "document.webp"),
              ("txt", "text", "document.txt"), ("mht", "html", "document.html")],
        )

    def test_bucket_migration_queues_pdf_to_dataset_when_staging_is_disabled(self):
        record = {"Repo": "VoiceOfML/Test", "File": "Table", "Extension": "xlsx",
                  "Folder": [], "Size": 100}
        queue = scan_reader_assets.build_queue(
            [record], self.revisions, reader_assets.empty_manifest(),
            bucket_migrate=True, bucket_pdf_staging=False,
        )
        self.assertEqual(len(queue), 1)
        self.assertEqual(queue[0]["reader_mode"], "pdf")
        self.assertTrue(queue[0]["page_stream"])
        self.assertNotIn("bucket_staging", queue[0])

    def test_bucket_migration_skips_an_asset_already_in_shared_bucket(self):
        key = reader_assets.asset_key("VoiceOfML/Test", "A/Book.docx")
        manifest = {"version": 1, "files": {key: {
            "status": "ready", "profile": "docx-native-v2", "bucket": reader_assets.READER_ASSETS_BUCKET,
        }}}
        self.assertEqual(scan_reader_assets.build_queue(
            self.records[:1], self.revisions, manifest, bucket_migrate=True,
        ), [])

    def test_bucket_migration_force_requeues_an_asset_already_marked_in_bucket(self):
        key = reader_assets.asset_key("VoiceOfML/Test", "A/Book.docx")
        manifest = {"version": 1, "files": {key: {
            "status": "ready", "profile": "docx-native-v2", "bucket": reader_assets.READER_ASSETS_BUCKET,
        }}}
        queue = scan_reader_assets.build_queue(
            self.records[:1], self.revisions, manifest, bucket_migrate=True, force=True,
        )
        self.assertEqual([item["key"] for item in queue], [key])

    def test_bucket_migration_repairs_legacy_chapter_path_without_force_rebuild(self):
        record = {"Repo": "VoiceOfML/Test", "File": "Book", "Extension": "epub",
                  "Folder": [], "Size": 100}
        key = reader_assets.asset_key(record["Repo"], "Book.epub")
        manifest = {"version": 1, "files": {key: {
            "status": "ready", "profile": "foliate-original-v1", "reader_mode": "foliate",
            "bucket": reader_assets.READER_ASSETS_BUCKET,
            "path": "objects/aa/document.epub",
            "chapter_manifest": "objects/aa/epub-chapters/chapter-manifest.json",
            "chapter_bucket": reader_assets.READER_ASSETS_BUCKET,
        }}}
        queue = scan_reader_assets.build_queue(
            [record], self.revisions, manifest, bucket_migrate=True,
        )
        self.assertEqual([item["key"] for item in queue], [key])
        self.assertFalse(queue[0].get("force_rebuild", False))

    def test_large_epub_without_chapters_is_requeued_for_upgrade(self):
        records = [{
            "Repo": "VoiceOfML/Test", "File": "Big", "Extension": "epub", "Folder": [],
            "Size": 1,
        }]
        revisions = {"VoiceOfML/Test": "rev1"}
        ready = {"status": "ready", "profile": "foliate-original-v1", "reader_mode": "foliate"}
        manifest = reader_assets.empty_manifest()
        manifest["files"] = {"VoiceOfML/Test\0Big.epub": dict(ready)}
        queue = scan_reader_assets.build_queue(records, revisions, manifest)
        self.assertEqual([item["path"] for item in queue], ["Big.epub"])
        manifest["files"]["VoiceOfML/Test\0Big.epub"]["chapter_bundle_profile"] = reader_assets.EPUB_CHAPTER_PROFILE
        manifest["files"]["VoiceOfML/Test\0Big.epub"]["chapter_manifest"] = "objects/a/chapter-manifest.json"
        self.assertEqual(scan_reader_assets.build_queue(records, revisions, manifest), [])

    def test_small_epub_and_mobi_are_queued_for_chapter_stream_upgrade(self):
        revisions = {"VoiceOfML/Test": "rev1"}
        manifest = reader_assets.empty_manifest()
        manifest["files"] = {
            "VoiceOfML/Test\0Small.epub": {"status": "ready", "profile": "foliate-original-v1"},
            "VoiceOfML/Test\0Small.mobi": {"status": "ready", "profile": "foliate-original-v1"},
        }
        records = [
            {"Repo": "VoiceOfML/Test", "File": "Small", "Extension": "epub", "Folder": [], "Size": 10},
            {"Repo": "VoiceOfML/Test", "File": "Small", "Extension": "mobi", "Folder": [], "Size": 10},
        ]
        self.assertEqual([item["extension"] for item in scan_reader_assets.build_queue(
            records, revisions, manifest,
        )], ["epub", "mobi"])
        for key in manifest["files"]:
            manifest["files"][key].update({
                "chapter_bundle_profile": reader_assets.EPUB_CHAPTER_PROFILE,
                "chapter_manifest": "ebook-chapters/objects/aa/chapter-manifest.json",
            })
        self.assertEqual(scan_reader_assets.build_queue(records, revisions, manifest), [])

    def test_html_resource_fragments_are_not_reader_documents(self):
        records = [{
            "Repo": "VoiceOfML/Test", "File": "script", "Folder": ["page_files"],
            "Extension": "htm", "Size": 10,
        }]
        self.assertEqual(
            scan_reader_assets.build_queue(records, self.revisions, reader_assets.empty_manifest()), []
        )

    def test_ready_current_profile_is_skipped(self):
        key = reader_assets.asset_key("VoiceOfML/Test", "A/Book.docx")
        manifest = {"version": 1, "files": {key: {
            "status": "ready", "source_revision": "rev1", "profile": "docx-native-v2"
        }}}
        self.assertEqual(scan_reader_assets.build_queue(self.records, self.revisions, manifest), [])

    def test_ready_manual_profile_is_skipped_for_any_source_revision(self):
        key = reader_assets.asset_key("VoiceOfML/Test", "A/Book.docx")
        manifest = {"version": 1, "files": {key: {
            "status": "ready", "source_revision": "rev1", "profile": "manual-pdf-v1"
        }}}
        self.assertEqual(scan_reader_assets.build_queue(self.records, self.revisions, manifest), [])
        self.assertEqual(scan_reader_assets.build_queue(
            self.records, {"VoiceOfML/Test": "rev2"}, manifest,
        ), [])

    def test_failed_current_profile_requires_retry_flag(self):
        key = reader_assets.asset_key("VoiceOfML/Test", "A/Book.docx")
        manifest = {"version": 1, "files": {key: {
            "status": "failed", "source_revision": "rev1", "profile": "docx-native-v2"
        }}}
        self.assertEqual(scan_reader_assets.build_queue(self.records, self.revisions, manifest), [])
        self.assertEqual(
            len(scan_reader_assets.build_queue(self.records, self.revisions, manifest, retry_failed=True)), 1
        )

    def test_failed_update_attempt_on_ready_asset_requires_retry_flag(self):
        key = reader_assets.asset_key("VoiceOfML/Test", "A/Book.docx")
        manifest = {"version": 1, "files": {key: {
            "status": "ready", "source_revision": "old", "profile": "docx-native-v2",
            "path": "objects/old/document.docx", "failed_source_revision": "rev1",
            "failed_profile": "docx-native-v2",
        }}}
        self.assertEqual(scan_reader_assets.build_queue(self.records, self.revisions, manifest), [])
        self.assertEqual(len(scan_reader_assets.build_queue(
            self.records, self.revisions, manifest, retry_failed=True,
        )), 1)

    def test_ready_fallback_with_failed_upgrade_is_not_retried_automatically(self):
        key = reader_assets.asset_key("VoiceOfML/Test", "A/Book.docx")
        manifest = {"version": 1, "files": {key: {
            "status": "ready", "source_revision": "old", "profile": "old-profile",
            "path": "objects/old/document.docx", "failed_source_revision": "rev1",
            "failed_profile": "docx-native-v2",
        }}}
        self.assertEqual(scan_reader_assets.build_queue(self.records, self.revisions, manifest), [])
        self.assertEqual(len(scan_reader_assets.build_queue(
            self.records, self.revisions, manifest, retry_failed=True,
        )), 1)

    def test_filter_and_limit_are_deterministic(self):
        records = self.records + [
            {"Repo": "VoiceOfML/Test", "File": "A", "Extension": "mobi", "Folder": [], "Size": 1},
            {"Repo": "VoiceOfML/Other", "File": "B", "Extension": "mobi", "Folder": [], "Size": 1},
        ]
        revisions = {**self.revisions, "VoiceOfML/Other": "rev2"}
        queue = scan_reader_assets.build_queue(
            records, revisions, reader_assets.empty_manifest(),
            repo="VoiceOfML/Test", extension="mobi", limit=1,
        )
        self.assertEqual([item["path"] for item in queue], ["A.mobi"])

    def test_stable_shards_are_disjoint_and_cover_the_queue(self):
        records = [
            {"Repo": "VoiceOfML/Test", "File": f"Book-{index}", "Extension": "docx", "Folder": [], "Size": 1}
            for index in range(40)
        ]
        complete = scan_reader_assets.build_queue(records, self.revisions, reader_assets.empty_manifest())
        shards = [
            scan_reader_assets.build_queue(
                records, self.revisions, reader_assets.empty_manifest(), shard_count=10, shard_index=index,
            )
            for index in range(10)
        ]
        shard_keys = [{item["key"] for item in shard} for shard in shards]
        self.assertEqual(set().union(*shard_keys), {item["key"] for item in complete})
        self.assertEqual(sum(len(keys) for keys in shard_keys), len(set().union(*shard_keys)))
        for index, shard in enumerate(shards):
            self.assertTrue(all(scan_reader_assets.shard_for_key(item["key"], 10) == index for item in shard))

    def test_shard_limit_is_applied_after_partitioning(self):
        records = [
            {"Repo": "VoiceOfML/Test", "File": f"Book-{index}", "Extension": "docx", "Folder": [], "Size": 1}
            for index in range(100)
        ]
        queue = scan_reader_assets.build_queue(
            records, self.revisions, reader_assets.empty_manifest(), limit=2, shard_count=4, shard_index=2,
        )
        self.assertEqual(len(queue), 2)
        self.assertTrue(all(scan_reader_assets.shard_for_key(item["key"], 4) == 2 for item in queue))

    def test_invalid_shard_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "invalid reader asset shard"):
            scan_reader_assets.build_queue(
                self.records, self.revisions, reader_assets.empty_manifest(), shard_count=2, shard_index=2,
            )

    def test_queues_only_password_marked_pdfs_without_exposing_password(self):
        records = self.records + [
            {"Repo": "VoiceOfML/Test", "File": "Protected（密码123）", "Extension": "pdf", "Folder": [], "Size": 1},
            {"Repo": "VoiceOfML/Test", "File": "Ordinary", "Extension": "pdf", "Folder": [], "Size": 1},
        ]
        queue = scan_reader_assets.build_queue(records, self.revisions, reader_assets.empty_manifest(), extension="pdf")
        self.assertEqual([item["path"] for item in queue], ["Protected（密码123）.pdf"])
        self.assertNotIn("source_password", queue[0])
        self.assertEqual(queue[0]["profile"], "qpdf-decrypted-v1")

    def test_unfiltered_queue_prioritizes_mature_bulk_formats(self):
        records = [
            {"Repo": "VoiceOfML/Test", "File": "A", "Extension": "doc", "Folder": [], "Size": 1},
            {"Repo": "VoiceOfML/Test", "File": "B", "Extension": "mobi", "Folder": [], "Size": 1},
            {"Repo": "VoiceOfML/Test", "File": "C", "Extension": "tif", "Folder": [], "Size": 1},
            {"Repo": "VoiceOfML/Test", "File": "D", "Extension": "djvu", "Folder": [], "Size": 1},
            {"Repo": "VoiceOfML/Test", "File": "E 密码123", "Extension": "pdf", "Folder": [], "Size": 1},
        ]
        queue = scan_reader_assets.build_queue(
            records, self.revisions, reader_assets.empty_manifest(),
        )
        self.assertEqual([item["extension"] for item in queue], ["tif", "mobi", "djvu", "doc", "pdf"])

    def test_exact_path_filter_selects_one_asset(self):
        queue = scan_reader_assets.build_queue(
            self.records, self.revisions, reader_assets.empty_manifest(), exact_path="A/Book.docx",
        )
        self.assertEqual([item["path"] for item in queue], ["A/Book.docx"])

    def test_missing_assets_repository_is_an_empty_manifest(self):
        api = Mock()
        response = requests.Response()
        response.status_code = 404
        response.request = requests.Request("GET", "https://huggingface.co/api/datasets/vomebook/Missing").prepare()
        api.file_exists.side_effect = RepositoryNotFoundError("missing", response=response)
        self.assertEqual(scan_reader_assets.remote_manifest(api, "vomebook/Missing"), reader_assets.empty_manifest())

    def test_reusable_objects_include_ready_files_and_orphans(self):
        manifest = {"version": 1, "files": {
            "book": {"status": "ready", "source_sha256": "a" * 64, "profile": "p1",
                     "path": "objects/a", "sha256": "b" * 64, "bytes": 10, "reader_mode": "pdf"},
        }, "orphans": {
            "objects/b": {"source_sha256": "c" * 64, "profile": "p2", "path": "objects/b",
                          "sha256": "d" * 64, "bytes": 20, "reader_mode": "epub", "since": "2026-01-01"},
        }}
        objects = scan_reader_assets.reusable_objects(manifest)
        self.assertEqual(objects["a" * 64 + "\0p1"]["path"], "objects/a")
        self.assertEqual(objects["c" * 64 + "\0p2"]["path"], "objects/b")

    def test_active_keys_can_be_empty_after_all_convertible_files_are_deleted(self):
        self.assertEqual(scan_reader_assets.active_keys([]), [])

    def test_active_keys_do_not_depend_on_revision_availability(self):
        self.assertEqual(
            scan_reader_assets.active_keys(self.records),
            [reader_assets.asset_key("VoiceOfML/Test", "A/Book.docx")],
        )

class ConverterTests(unittest.TestCase):
    def test_download_source_resumes_partial_response(self):
        class Response:
            status = 206
            headers = {"Content-Range": "bytes 6-10/11"}
            def __enter__(self):
                return self
            def __exit__(self, *_):
                return False
            def read(self, _size):
                if self.done:
                    return b""
                self.done = True
                return b"world"

        with tempfile.TemporaryDirectory() as root:
            target = Path(root) / "source.bin"
            target.write_bytes(b"hello ")
            response = Response()
            response.done = False
            with patch.object(urllib.request, "urlopen", return_value=response) as open_url:
                digest, size = convert_reader_assets.download_source("https://example.test/source", target)
            self.assertEqual(target.read_bytes(), b"hello world")
            self.assertEqual(size, 11)
            self.assertEqual(digest, hashlib.sha256(b"hello world").hexdigest())
            self.assertEqual(open_url.call_args.args[0].headers["Range"], "bytes=6-")

    def test_download_source_rejects_invalid_content_range(self):
        class Response:
            status = 206
            headers = {"Content-Range": "bytes 0-4/11"}
            def __enter__(self): return self
            def __exit__(self, *_): return False
            def read(self, _size): return b"wrong"

        with tempfile.TemporaryDirectory() as root:
            target = Path(root) / "source.bin"
            target.write_bytes(b"hello ")
            with patch.object(urllib.request, "urlopen", return_value=Response()), \
                    patch.object(convert_reader_assets.time, "sleep"):
                with self.assertRaises(http.client.IncompleteRead):
                    convert_reader_assets.download_source("https://example.test/source", target)
            self.assertFalse(target.exists())

    def test_tiff_conversion_preserves_multiple_frames(self):
        from PIL import Image

        with tempfile.TemporaryDirectory() as root:
            source, target = Path(root) / "source.tiff", Path(root) / "document.pdf"
            first = Image.new("RGB", (20, 30), "white")
            second = Image.new("RGB", (20, 30), "black")
            first.save(source, save_all=True, append_images=[second])
            def merge(command):
                Path(command[-1]).write_bytes(b"%PDF-merged")

            with patch.object(convert_reader_assets, "run_checked", side_effect=merge) as run:
                convert_reader_assets.convert_tiff(source, target, Path(root))
            self.assertEqual(target.read_bytes()[:5], b"%PDF-")
            self.assertGreater(target.stat().st_size, 0)
            self.assertEqual(run.call_args.args[0][0], "pdfunite")
            self.assertEqual(len(run.call_args.args[0][1:-1]), 2)

    def test_single_frame_tiff_does_not_require_pdfunite(self):
        from PIL import Image

        with tempfile.TemporaryDirectory() as root:
            source, target = Path(root) / "source.tif", Path(root) / "document.pdf"
            Image.new("RGB", (20, 30), "white").save(source)
            with patch.object(convert_reader_assets, "run_checked") as run:
                convert_reader_assets.convert_tiff(source, target, Path(root))
            run.assert_not_called()
            self.assertEqual(target.read_bytes()[:5], b"%PDF-")

    def test_djvu_conversion_uses_ddjvu_pdf_output(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.djvu", work / "document.pdf"
            source.write_bytes(b"DJVU")
            with patch.object(convert_reader_assets, "run_checked") as run:
                convert_reader_assets.convert_file({"extension": "djvu"}, source, target, work)
            self.assertEqual(run.call_args.args[0], [
                "ddjvu", "-format=pdf", str(source), str(target),
            ])
            self.assertEqual(run.call_args.kwargs["timeout_seconds"], convert_reader_assets.DJVU_COMMAND_TIMEOUT_SECONDS)

    def test_djvu_conversion_retries_command_once(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.djvu", work / "document.pdf"
            source.write_bytes(b"DJVU")
            with patch.object(convert_reader_assets, "run_checked", side_effect=[
                    convert_reader_assets.ReaderConversionCommandError("temporary"), None,
            ]) as run, patch.object(convert_reader_assets.time, "sleep"):
                convert_reader_assets.convert_file({"extension": "djvu"}, source, target, work)
            self.assertEqual(run.call_count, 2)

    def test_audio_conversion_uses_bounded_mp3_contract(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.ape", work / "audio.mp3"
            source.write_bytes(b"audio")
            with patch.object(convert_reader_assets, "run_checked") as run:
                convert_reader_assets.convert_file({"extension": "ape"}, source, target, work)
            command = run.call_args.args[0]
            self.assertEqual(command[0:6], ["ffmpeg", "-nostdin", "-hide_banner", "-loglevel", "error", "-y"])
            self.assertIn("libmp3lame", command)
            self.assertIn("0:a:0", command)
            self.assertEqual(run.call_args.kwargs["timeout_seconds"], convert_reader_assets.MEDIA_COMMAND_TIMEOUT_SECONDS)

    def test_video_conversion_uses_h264_aac_faststart_contract(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.flv", work / "video.mp4"
            source.write_bytes(b"video")
            with patch.object(convert_reader_assets, "run_checked") as run:
                convert_reader_assets.convert_file({"extension": "flv"}, source, target, work)
            command = run.call_args.args[0]
            for value in ("libx264", "yuv420p", "aac", "+faststart", "0:v:0", "0:a:0?"):
                self.assertIn(value, command)
            self.assertEqual(run.call_args.kwargs["timeout_seconds"], convert_reader_assets.MEDIA_COMMAND_TIMEOUT_SECONDS)

    def test_audio_only_rm_uses_black_video_and_aac_audio(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.rm", work / "video.mp4"
            source.write_bytes(b"rm")
            item = {"extension": "rm", "source_media_mode": "audio"}
            with patch.object(convert_reader_assets, "run_checked") as run:
                convert_reader_assets.convert_file(item, source, target, work)
            command = run.call_args.args[0]
            self.assertIn("color=c=black:s=640x360:r=1", command)
            self.assertIn("0:v:0", command)
            self.assertIn("1:a:0", command)
            self.assertIn("libx264", command)
            self.assertIn("aac", command)

    def test_static_image_conversion_outputs_webp(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.png", work / "document.webp"
            convert_reader_assets.Image.new("RGBA", (4, 4), (255, 0, 0, 128)).save(source)
            convert_reader_assets.convert_file(
                {"extension": "png", "reader_mode": "image", "output_name": "document.webp"},
                source, target, work,
            )
            self.assertEqual(target.read_bytes()[:4], b"RIFF")
            with convert_reader_assets.Image.open(target) as image:
                self.assertEqual(image.format, "WEBP")

    def test_calibre_office_book_conversion_uses_html_output(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.odt", work / "document.html"
            source.write_bytes(b"office book")
            def fake_run(command, **_kwargs):
                if command[0] == "libreoffice":
                    (work / "odt-html" / "source.html").write_text("<p>body</p>", encoding="utf-8")
            with patch.object(convert_reader_assets, "run_checked", side_effect=fake_run) as run:
                convert_reader_assets.convert_file({"extension": "odt"}, source, target, work)
            self.assertEqual(run.call_args.args[0][0:2], ["libreoffice", "--headless"])

    def test_rtf_falls_back_to_libreoffice_html_before_calibre(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.rtf", work / "document.html"
            source.write_bytes(b"rtf")
            calls = []
            def fake_run(command, **_kwargs):
                calls.append(command)
                if command[0] == "ebook-convert" and len(calls) == 1:
                    raise convert_reader_assets.ReaderConversionCommandError("unsupported codepage")
                if command[0] == "libreoffice":
                    (work / "rtf-html" / "source.html").write_text("<p>body</p>", encoding="utf-8")
            with patch.object(convert_reader_assets, "download_source", return_value=("a" * 64, 100)), patch.object(
                    convert_reader_assets, "run_checked", side_effect=fake_run):
                convert_reader_assets.convert_file({"extension": "rtf"}, source, target, work)
            self.assertEqual([call[0] for call in calls], ["libreoffice"])
            self.assertEqual(calls[0][calls[0].index("--convert-to") + 1], "html")

    def test_ps_content_validation_renders_sample_pages(self):
        with patch.object(convert_reader_assets, "validate_pdf_content") as validate:
            convert_reader_assets.validate_reader_content(
                Path("document.pdf"), {"extension": "ps", "reader_mode": "pdf"}, Path("work"),
            )
        validate.assert_called_once_with(Path("document.pdf"), Path("work"))

    def test_office_presentation_conversion_uses_libreoffice_pdf(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.pptx", work / "document.pdf"
            source.write_bytes(b"slides")
            def fake_run(command, **_kwargs):
                if command[0] == "libreoffice":
                    (work / "office-pdf" / "source.pdf").write_bytes(b"%PDF-office")
            with patch.object(convert_reader_assets, "run_checked", side_effect=fake_run) as run:
                convert_reader_assets.convert_file({"extension": "pptx"}, source, target, work)
            self.assertEqual(run.call_args.args[0][0:2], ["libreoffice", "--headless"])
            self.assertEqual(run.call_args.args[0][run.call_args.args[0].index("--convert-to") + 1], "pdf")
            self.assertEqual(target.read_bytes(), b"%PDF-office")

    def test_mislabeled_ole_xlsx_uses_xls_extension(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.xlsx", work / "document.pdf"
            source.write_bytes(convert_reader_assets.OLE_SIGNATURE + b"spreadsheet")
            def fake_run(command, **_kwargs):
                (work / "office-pdf" / "source.pdf").write_bytes(b"%PDF-sheet")
            with patch.object(convert_reader_assets, "run_checked", side_effect=fake_run) as run:
                convert_reader_assets.convert_file({"extension": "xlsx"}, source, target, work)
            self.assertTrue(str(run.call_args.args[0][-1]).endswith("source.xls"))
            self.assertEqual(target.read_bytes(), b"%PDF-sheet")

    def test_mislabeled_ole_xlsx_ignores_inapplicable_workflow_password(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.xlsx", work / "document.pdf"
            source.write_bytes(convert_reader_assets.OLE_SIGNATURE + b"spreadsheet")

            def fake_run(command, **_kwargs):
                (work / "office-pdf" / "source.pdf").write_bytes(b"%PDF-sheet")

            with patch.dict(convert_reader_assets.os.environ, {"READER_CONVERSION_PASSWORD": "secret"}), \
                    patch.dict("sys.modules", {"msoffcrypto": Mock(OfficeFile=Mock(side_effect=ValueError("not encrypted")))}), \
                    patch.object(convert_reader_assets, "run_checked", side_effect=fake_run) as run:
                convert_reader_assets.convert_file({"extension": "xlsx", "path": "book.xlsx"}, source, target, work)
            self.assertTrue(str(run.call_args.args[0][-1]).endswith("source.xls"))

    def test_encrypted_ole_xlsx_uses_temporary_conversion_password(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.xlsx", work / "document.pdf"
            source.write_bytes(convert_reader_assets.OLE_SIGNATURE + b"encrypted")
            office = Mock()
            source_handles = []

            def office_file(handle):
                source_handles.append(handle)
                return office

            def is_encrypted():
                self.assertFalse(source_handles[0].closed)
                return True

            office.is_encrypted.side_effect = is_encrypted

            def decrypt(target_file):
                target_file.write(b"decrypted")

            office.decrypt.side_effect = decrypt
            with patch.dict(convert_reader_assets.os.environ, {"READER_CONVERSION_PASSWORD": "secret"}), \
                    patch.dict("sys.modules", {"msoffcrypto": Mock(OfficeFile=office_file)}), \
                    patch.object(convert_reader_assets, "run_checked") as run:
                def fake_run(command, **_kwargs):
                    (work / "office-pdf" / "decrypted.pdf").write_bytes(b"%PDF-sheet")
                run.side_effect = fake_run
                convert_reader_assets.convert_file({"extension": "xlsx"}, source, target, work)
            self.assertTrue(str(run.call_args.args[0][-1]).endswith("decrypted.xlsx"))

    def test_mht_conversion_reuses_mhtml_sanitizer(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.mht", work / "document.html"
            source.write_bytes(b"MIME-Version: 1.0")
            with patch.object(convert_reader_assets, "mhtml_to_html") as convert:
                convert_reader_assets.convert_file({"extension": "mht"}, source, target, work)
            convert.assert_called_once_with(source, target)

    def test_mhtml_windows_content_location_does_not_break_resource_joining(self):
        message = b"""MIME-Version: 1.0
Content-Type: multipart/related; boundary=x

--x
Content-Type: text/html; charset=utf-8
Content-Location: file://D:\\books\\page.htm

<html><body><p>Readable Windows MHTML body</p><img src="image.png"></body></html>
--x
Content-Type: image/png
Content-Transfer-Encoding: base64
Content-Location: file://D:\\books\\image.png

aW1hZ2U=
--x--
"""
        with tempfile.TemporaryDirectory() as root:
            source, target = Path(root) / "source.mht", Path(root) / "document.html"
            source.write_bytes(message)
            convert_reader_assets.mhtml_to_html(source, target)
            self.assertIn("Readable Windows MHTML body", target.read_text(encoding="utf-8"))

    def test_ps_conversion_uses_ghostscript_pdf(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.ps", work / "document.pdf"
            source.write_bytes(b"%!PS")
            with patch.object(convert_reader_assets, "run_checked") as run:
                convert_reader_assets.convert_file({"extension": "ps"}, source, target, work)
            command = run.call_args.args[0]
            self.assertIn("-sDEVICE=pdfwrite", command)
            self.assertIn(f"-sOutputFile={target}", command)
            self.assertEqual(run.call_args.kwargs["timeout_seconds"], convert_reader_assets.POSTSCRIPT_COMMAND_TIMEOUT_SECONDS)

    def test_media_validation_requires_browser_compatible_streams(self):
        audio = {"format": {"format_name": "mp3", "duration": "10"}, "streams": [
            {"codec_type": "audio", "codec_name": "mp3"},
        ]}
        video = {"format": {"format_name": "mov,mp4,m4a,3gp,3g2,mj2", "duration": "20"}, "streams": [
            {"codec_type": "video", "codec_name": "h264", "pix_fmt": "yuv420p", "width": 1920, "height": 1080},
            {"codec_type": "audio", "codec_name": "aac"},
        ]}
        with patch.object(convert_reader_assets, "media_probe", return_value=audio):
            convert_reader_assets.validate_media_output(Path("audio.mp3"), "audio")
        with patch.object(convert_reader_assets, "media_probe", return_value=video):
            convert_reader_assets.validate_media_output(Path("video.mp4"), "video")
        video["streams"][0]["codec_name"] = "hevc"
        with patch.object(convert_reader_assets, "media_probe", return_value=video), self.assertRaises(RuntimeError):
            convert_reader_assets.validate_media_output(Path("video.mp4"), "video")

    def test_mobi_is_copied_for_foliate(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.mobi", work / "document.mobi"
            source.write_bytes(b"BOOKMOBI")
            convert_reader_assets.convert_file({"extension": "mobi", "reader_mode": "foliate"}, source, target, work)
            self.assertEqual(target.read_bytes(), b"BOOKMOBI")

    def test_azw3_is_copied_for_foliate(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.azw3", work / "document.azw3"
            source.write_bytes(b"BOOKAZW3")
            convert_reader_assets.convert_file({"extension": "azw3", "reader_mode": "foliate"}, source, target, work)
            self.assertEqual(target.read_bytes(), b"BOOKAZW3")

    def test_chm_conversion_uses_native_epub_output(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.chm", work / "document.epub"
            source.write_bytes(b"ITSF")
            with patch.object(convert_reader_assets, "convert_chm") as conversion:
                convert_reader_assets.convert_file({"extension": "chm"}, source, target, work)
            conversion.assert_called_once_with(source, target, work)

    def test_caj_family_detection_uses_content_not_extension(self):
        cases = {
            b"%PDF-1.7": "pdf", b"KDH 2.00": "kdh", b"CAJ\0": "caj",
            b"HN\0\0": "hn", b"\xc8\0\0\0": "c8", b"\0" * 16: "hn",
        }
        with tempfile.TemporaryDirectory() as root:
            source = Path(root) / "source.caj"
            for header, expected in cases.items():
                source.write_bytes(header)
                self.assertEqual(convert_reader_assets.detect_caj_family(source), expected)

    def test_caj_family_copies_embedded_pdf_and_dispatches_converter(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.caj", work / "document.pdf"
            source.write_bytes(b"%PDF-1.7\nbody")
            convert_reader_assets.convert_caj_family(source, target, work)
            self.assertEqual(target.read_bytes(), source.read_bytes())

            source.write_bytes(b"CAJ\0data")
            converter = work / "caj2pdf" / "caj2pdf"
            converter.parent.mkdir()
            converter.write_text("converter", encoding="utf-8")
            for library in ("libjbigdec.so", "libjbig2codec.so"):
                (converter.parent / library).write_bytes(library.encode())
            with patch.object(convert_reader_assets, "CAJ2PDF_DIR", converter.parent), \
                    patch.object(convert_reader_assets, "run_checked") as run:
                convert_reader_assets.convert_caj_family(source, target, work)
            self.assertEqual(run.call_args.args[0], ["python3", str(converter), "convert", str(source), "--output", str(target)])
            self.assertEqual(run.call_args.kwargs["cwd"], work)
            self.assertEqual((work / "libjbigdec.so").read_bytes(), b"libjbigdec.so")
            self.assertEqual((work / "libjbig2codec.so").read_bytes(), b"libjbig2codec.so")

    def test_kdh_extraction_decrypts_embedded_pdf_before_repair(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.kdh", work / "document.pdf"
            payload = b"%PDF-1.3\nbody\n%%EOFtrailer"
            key = b"FZHMEI"
            encrypted = bytes(value ^ key[index % len(key)] for index, value in enumerate(payload))
            source.write_bytes(b"KDH " + b"\0" * 250 + encrypted)

            def repair(command):
                self.assertEqual(Path(command[2]).read_bytes(), b"%PDF-1.3\nbody\n%%EOF")
                Path(command[3]).write_bytes(Path(command[2]).read_bytes())

            with patch.object(convert_reader_assets, "run_checked", side_effect=repair):
                convert_reader_assets.extract_kdh_pdf(source, target, work)
            self.assertTrue(target.read_bytes().startswith(b"%PDF-"))

    def test_mhtml_fallback_extracts_html_and_inlines_resources(self):
        message = b"""MIME-Version: 1.0
Content-Type: multipart/related; boundary=x

--x
Content-Type: text/html; charset=utf-8
Content-Location: file:///book/index.html

<html><head><base href="https://example.test/"></head><body onload="alert(1)"><p>Readable CHM body text</p><script>alert(1)</script><object>unsafe</object><img src="image.png"><a href="javascript:alert(1)">bad</a></body></html>
--x
Content-Type: image/png
Content-Transfer-Encoding: base64
Content-Location: file:///book/image.png

aW1hZ2U=
--x--
"""
        with tempfile.TemporaryDirectory() as root:
            source, target = Path(root) / "book.mht", Path(root) / "book.html"
            source.write_bytes(message)
            convert_reader_assets.mhtml_to_html(source, target)
            document = target.read_text(encoding="utf-8")
            self.assertIn("Readable CHM body text", document)
            self.assertIn("data:image/png;base64,aW1hZ2U=", document)
            self.assertNotRegex(document, r"(?i)<(?:base|object|script)\b|onload=|javascript:")

    def test_chm_conversion_falls_back_to_embedded_mhtml_for_empty_epub(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.chm", work / "book.epub"
            source.write_bytes(b"ITSF")
            calls = []

            def run(command, **_kwargs):
                calls.append(command)

            with patch.object(convert_reader_assets, "run_checked", side_effect=run), \
                    patch.object(convert_reader_assets, "convert_chm_to_html", side_effect=lambda _source, output, _work: output.write_text("<p>body</p>", encoding="utf-8")), \
                    patch.object(convert_reader_assets, "sanitize_chm_epub"), \
                    patch.object(convert_reader_assets, "validate_output"), \
                    patch.object(convert_reader_assets, "validate_chm_epub", side_effect=[RuntimeError("converted CHM EPUB has no readable content"), None]):
                convert_reader_assets.convert_chm(source, target, work)
            self.assertEqual([call[0] for call in calls], ["ebook-convert", "ebook-convert"])

    def test_chm_conversion_falls_back_to_direct_html_pages(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.chm", work / "book.epub"
            source.write_bytes(b"ITSF")
            calls = []

            def run(command, **_kwargs):
                calls.append(command)

            with patch.object(convert_reader_assets, "run_checked", side_effect=run), \
                    patch.object(convert_reader_assets, "convert_chm_to_html", side_effect=lambda _source, output, _work: output.write_text("<p>第一章</p><p>第二章</p>", encoding="utf-8")), \
                    patch.object(convert_reader_assets, "sanitize_chm_epub"), \
                    patch.object(convert_reader_assets, "validate_output"), \
                    patch.object(convert_reader_assets, "validate_chm_epub", side_effect=[RuntimeError("empty"), None]), \
                    patch.object(convert_reader_assets, "validate_reader_content"):
                convert_reader_assets.convert_chm(source, target, work)
            self.assertEqual([call[0] for call in calls], ["ebook-convert", "ebook-convert"])
            self.assertIn("第一章", (work / "chm-fallback.html").read_text(encoding="utf-8"))

    def test_chm_epub_sanitizer_removes_active_and_external_content(self):
        with tempfile.TemporaryDirectory() as root:
            work, epub = Path(root), Path(root) / "book.epub"
            with zipfile.ZipFile(epub, "w") as archive:
                archive.writestr("mimetype", "application/epub+zip")
                archive.writestr("chapter.xhtml", '<html xmlns="http://www.w3.org/1999/xhtml"><head><base href="https://example.test/"/></head><body onload="bad()"><script>bad()</script><object>bad</object><img src="https://example.test/a.png"/><a href="javascript:bad()">bad</a><p>正文内容保持不变。</p></body></html>')
                archive.writestr("style.css", '@import "https://example.test/a.css"; body{background:url(//example.test/a.png)}')
            convert_reader_assets.sanitize_chm_epub(epub, work)
            with zipfile.ZipFile(epub) as archive:
                content = archive.read("chapter.xhtml").decode() + archive.read("style.css").decode()
                self.assertEqual(archive.infolist()[0].filename, "mimetype")
                self.assertEqual(archive.getinfo("mimetype").compress_type, zipfile.ZIP_STORED)
            self.assertNotRegex(content, r"(?i)<(?:base|object|script)\b|onload=|javascript:|https://example|//example")
            self.assertIn("正文内容保持不变", content)

    def test_epub_sanitizer_preserves_valid_xhtml_and_internal_resources(self):
        with tempfile.TemporaryDirectory() as root:
            work, epub = Path(root), Path(root) / "book.epub"
            with zipfile.ZipFile(epub, "w") as archive:
                archive.writestr("mimetype", "application/epub+zip")
                archive.writestr("META-INF/container.xml", '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>')
                archive.writestr("OEBPS/content.opf", '<package xmlns="http://www.idpf.org/2007/opf"><manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="chapter"/></spine></package>')
                archive.writestr("OEBPS/chapter.xhtml", '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Book</title><link rel="stylesheet" href="style.css"/></head><body><p>This is enough readable chapter text.</p><img src="images/cover.png" alt="cover"/></body></html>')
                archive.writestr("OEBPS/style.css", "p{color:red}")
                archive.writestr("OEBPS/images/cover.png", b"PNG")
            convert_reader_assets.sanitize_chm_epub(epub, work)
            convert_reader_assets.validate_epub_content(epub)
            with zipfile.ZipFile(epub) as archive:
                chapter = archive.read("OEBPS/chapter.xhtml").decode()
            self.assertIn("style.css", chapter)
            self.assertIn("images/cover.png", chapter)
            self.assertIn("<", chapter)

    def test_epub_html_extension_preserves_xml_document_structure(self):
        with tempfile.TemporaryDirectory() as root:
            work, epub = Path(root), Path(root) / "book.epub"
            with zipfile.ZipFile(epub, "w") as archive:
                archive.writestr("mimetype", "application/epub+zip")
                archive.writestr("chapter.html", '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Book</title></head><body><p>Readable chapter text.</p><img src="images/a.png"/></body></html>')
            convert_reader_assets.sanitize_chm_epub(epub, work)
            with zipfile.ZipFile(epub) as archive:
                chapter = archive.read("chapter.html").decode()
            parsed = convert_reader_assets.ET.fromstring(chapter)
            self.assertEqual(parsed.tag.rsplit("}", 1)[-1], "html")
            self.assertIn("images/a.png", chapter)

    def test_css_sanitizer_removes_external_urls(self):
        sanitized = convert_reader_assets.sanitize_css(
            "p{color:red;fill:url(https://evil.test/a.svg);cursor:url(//evil.test/a.cur)}"
        )
        self.assertIn("color:red", sanitized)
        self.assertNotIn("evil.test", sanitized)

    def test_chm_epub_validation_rejects_cover_without_body_text(self):
        with tempfile.TemporaryDirectory() as root:
            epub = Path(root) / "book.epub"
            with zipfile.ZipFile(epub, "w") as archive:
                archive.writestr("mimetype", "application/epub+zip")
                archive.writestr("META-INF/container.xml", """<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="content.opf"/></rootfiles></container>""")
                archive.writestr("content.opf", """<package xmlns="http://www.idpf.org/2007/opf"><manifest><item id="cover" href="cover.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="cover"/></spine></package>""")
                archive.writestr("cover.xhtml", "<html xmlns=\"http://www.w3.org/1999/xhtml\"><head><style>body has lots of fake readable style content</style></head><body><svg xmlns=\"http://www.w3.org/2000/svg\"/></body></html>")
            with self.assertRaisesRegex(RuntimeError, "no readable content"):
                convert_reader_assets.validate_chm_epub(epub)

    def test_chm_epub_validation_requires_readable_safe_spine(self):
        with tempfile.TemporaryDirectory() as root:
            epub = Path(root) / "book.epub"
            def write_epub(document):
                with zipfile.ZipFile(epub, "w") as archive:
                    archive.writestr("mimetype", "application/epub+zip")
                    archive.writestr("META-INF/container.xml", """<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>""")
                    archive.writestr("OEBPS/content.opf", """<package xmlns="http://www.idpf.org/2007/opf"><manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="chapter"/></spine></package>""")
                    archive.writestr("OEBPS/chapter.xhtml", document)

            write_epub("<html><body><p>这是可以正常阅读的中文 CHM 转换正文内容。</p></body></html>")
            convert_reader_assets.validate_chm_epub(epub)

            write_epub("<html><body><script>alert(1)</script></body></html>")
            with self.assertRaisesRegex(RuntimeError, "active content"):
                convert_reader_assets.validate_chm_epub(epub)

    def test_chm_document_write_literals_are_expanded_without_execution(self):
        document, expanded = convert_reader_assets.expand_document_writes(
            "document.write(\"<p>第一段</p>\" + '<p>第二段\\x21</p>');"
        )
        self.assertTrue(expanded)
        self.assertEqual(document, "<p>第一段</p><p>第二段!</p>")

    def test_chm_document_write_rejects_dynamic_expressions(self):
        document, expanded = convert_reader_assets.expand_document_writes(
            "document.write('<p>' + userInput + '</p>');"
        )
        self.assertFalse(expanded)
        self.assertEqual(document, "document.write('<p>' + userInput + '</p>');")

    def test_chm_static_writes_follow_javascript_escaping_and_skip_inactive_code(self):
        text, found = convert_reader_assets.expand_document_writes(
            '/* document.write("wrong") */ function unused(){document.write("wrong");}'
            'if(false){document.write("wrong");}'
            'document.write("<p>\\u4e2d\\x21<\\/p>"); document.write(userInput);'
            'document.writeln("next", " page");'
        )
        self.assertTrue(found)
        self.assertEqual(text, '<p>中!</p>next page\n')

    def test_xml_sanitizer_preserves_legacy_container_text_and_removed_element_tails(self):
        root = convert_reader_assets.ET.fromstring(convert_reader_assets.sanitize_xml_document(
            '<html><body>before<font color="red">text<b>bold</b>tail</font>after'
            '<script>discard()</script>keep<center>center</center>end</body></html>'
        ))
        self.assertEqual(''.join(root.find('body').itertext()), 'beforetextboldtailafterkeepcenterend')

    def test_small_epub_direct_copy_uses_foliate(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.epub", work / "document.epub"
            source.write_bytes(b"epub")
            convert_reader_assets.convert_file({"extension": "epub", "reader_mode": "foliate"}, source, target, work)
            self.assertEqual(target.read_bytes(), b"epub")

    def test_validation_rejects_malformed_outputs_for_each_active_container(self):
        with tempfile.TemporaryDirectory() as root:
            root = Path(root)
            bad_pdf = root / "bad.pdf"; bad_pdf.write_bytes(b"not pdf")
            with self.assertRaises(RuntimeError): convert_reader_assets.validate_output(bad_pdf, "pdf")
            bad_html = root / "bad.html"; bad_html.write_bytes(b"")
            with self.assertRaises(RuntimeError): convert_reader_assets.validate_output(bad_html, "html")
            bad_epub = root / "bad.epub"
            with zipfile.ZipFile(bad_epub, "w") as archive: archive.writestr("mimetype", "wrong")
            with self.assertRaises(RuntimeError): convert_reader_assets.validate_output(bad_epub, "epub")

    def test_djvu_pdf_validation_checks_first_and_last_pages(self):
        with tempfile.TemporaryDirectory() as root:
            work, pdf = Path(root), Path(root) / "document.pdf"
            pdf.write_bytes(b"%PDF-test")

            def render(command, **_kwargs):
                Path(command[-1]).with_suffix(".png").write_bytes(b"png")

            with patch.object(convert_reader_assets, "command_output", return_value="Pages:           12\n"), \
                    patch.object(convert_reader_assets, "run_checked", side_effect=render) as run:
                convert_reader_assets.validate_djvu_pdf(pdf, work)
            self.assertEqual([call.args[0][2] for call in run.call_args_list], ["1", "12"])

    def test_doc_conversion_uses_headless_libreoffice_to_docx(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.doc", work / "document.docx"
            source.write_bytes(b"doc")

            def fake_run(command):
                (work / "office" / "source.docx").write_bytes(b"docx")

            with patch.object(convert_reader_assets, "run_checked", side_effect=fake_run) as run:
                convert_reader_assets.convert_file({"extension": "doc"}, source, target, work)
            command = run.call_args.args[0]
            self.assertEqual(command[:2], ["libreoffice", "--headless"])
            self.assertTrue(any(arg.startswith("-env:UserInstallation=file://") for arg in command))
            self.assertEqual(command[command.index("--convert-to") + 1], "docx")
            self.assertEqual(target.read_bytes(), b"docx")

    def test_epub_conversion_does_not_linearize_pdf(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.epub", work / "document.epub"
            source.write_bytes(b"epub")
            convert_reader_assets.convert_file({"extension": "epub", "reader_mode": "foliate"}, source, target, work)
            self.assertEqual(target.read_bytes(), b"epub")

    def test_doc_and_docx_use_distinct_docx_profiles(self):
        self.assertEqual(reader_assets.conversion_contract("doc"), ("libreoffice-docx-v2", "docx", "document.docx"))
        self.assertEqual(reader_assets.conversion_contract("docx"), ("docx-native-v2", "docx", "document.docx"))

    def test_html_and_htm_use_sanitized_html_profile(self):
        expected = ("sanitized-html-v5", "html", "document.html")
        self.assertEqual(reader_assets.conversion_contract("html"), expected)
        self.assertEqual(reader_assets.conversion_contract("htm"), expected)

    def test_html_conversion_preserves_safe_source_html(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.html", work / "document.html"
            source.write_bytes(b"<html><body>text</body></html>")

            convert_reader_assets.convert_file({"extension": "html"}, source, target, work)
            self.assertEqual(target.read_text(encoding="utf-8"), "text")

    def test_generated_html_inlines_local_images_and_stylesheets(self):
        with tempfile.TemporaryDirectory() as root:
            root = Path(root)
            (root / "image.png").write_bytes(b"png")
            (root / "style.css").write_text("p { color: red; }", encoding="utf-8")
            result = convert_reader_assets.inline_local_html_resources(
                '<link rel="stylesheet" href="style.css"><p><img src="image.png">正文</p>',
                root, root,
            )
            self.assertIn("data:image/png;base64", result)
            self.assertIn("color: red", result)
            self.assertNotIn('href="style.css"', result)

    def test_generated_html_resolves_local_resources_case_insensitively(self):
        with tempfile.TemporaryDirectory() as root:
            root = Path(root)
            (root / "Images").mkdir()
            (root / "Images" / "screen.PNG").write_bytes(b"png")
            result = convert_reader_assets.inline_local_html_resources(
                '<p><img src="images/SCREEN.png">正文</p>', root, root,
            )
            self.assertIn("data:image/png;base64", result)

    def test_chm_epub_repair_adds_case_insensitive_source_images(self):
        with tempfile.TemporaryDirectory() as root:
            root = Path(root)
            source = root / "source.chm"
            epub = root / "book.epub"
            with zipfile.ZipFile(source, "w") as archive:
                archive.writestr("crack.htm", '<html><body><p>正文</p><img src="Images/screen.jpg"></body></html>')
                archive.writestr("images/screen.jpg", b"JPEG")
            with zipfile.ZipFile(epub, "w") as archive:
                archive.writestr("mimetype", "application/epub+zip")
                archive.writestr("META-INF/container.xml", '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>')
                archive.writestr("OEBPS/content.opf", '<package xmlns="http://www.idpf.org/2007/opf"><manifest><item id="chapter" href="crack.htm" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="chapter"/></spine></package>')
                archive.writestr("OEBPS/crack.htm", '<html xmlns="http://www.w3.org/1999/xhtml"><body><p>正文</p><img src="Images/screen.jpg"/></body></html>')
            report = convert_reader_assets.repair_chm_epub_images(source, epub, root / "work")
            self.assertEqual(len(report["added"]), 1)
            self.assertEqual(report["missing"], [])
            with zipfile.ZipFile(epub) as archive:
                self.assertEqual(archive.read("OEBPS/Images/screen.jpg"), b"JPEG")
                self.assertIn("chm-image-1", archive.read("OEBPS/content.opf").decode())

    def test_generated_html_expands_static_writes_and_drops_dynamic_scripts(self):
        result = convert_reader_assets.inline_local_html_resources(
            "<script>document.write('<p>生成正文</p>');</script>"
            "<script>document.write('<p>' + dynamicValue + '</p>');</script>",
            Path("/tmp"), Path("/tmp"),
        )
        self.assertIn("生成正文", result)
        self.assertNotIn("dynamicValue", result)
        self.assertNotIn("document.write", result)

    def test_html_conversion_removes_active_and_remote_content(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.html", work / "document.html"
            source.write_text(
                '<meta http-equiv="refresh" content="0;url=https://evil.test">'
                '<script>alert(1)</script><form action="https://evil.test"><p onclick="x()">safe</p></form>'
                '<a href="https://evil.test">remote</a><img src="//evil.test/x.png">',
                encoding="utf-8",
            )
            convert_reader_assets.convert_file(
                {"extension": "html", "source_url": "https://huggingface.co/datasets/VoiceOfML/Test/resolve/rev/source.html"},
                source, target, work,
            )
            text = target.read_text(encoding="utf-8")
            self.assertIn("safe", text)
            self.assertIn("remote", text)
            for unsafe in ("http-equiv", "<script", "<form", "onclick", "https://evil.test", "//evil.test"):
                self.assertNotIn(unsafe, text)

    def test_html_sanitizer_removes_unquoted_active_and_remote_urls(self):
        sanitized = convert_reader_assets.sanitize_html(
            "<script>alert(1)<img srcset='https://evil.test/a 1x'>"
            "<a href=javascript:alert(1)>x</a><img src=https://evil.test/x>"
            "<img src=data:text/html,evil><svg><a xlink:href='javascript:alert(2)'>x</a></svg>"
        )
        self.assertNotIn("javascript:", sanitized)
        self.assertNotIn("https://", sanitized)
        self.assertNotIn("data:text/html", sanitized)
        self.assertNotIn("srcset", sanitized)
        self.assertNotIn("<script", sanitized)
        self.assertNotIn("<svg", sanitized)

    def test_html_conversion_decodes_declared_legacy_charset(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.html", work / "document.html"
            source.write_bytes(
                '<meta http-equiv="Content-Type" content="text/html; charset=GB2312"><p>列宁</p>'.encode("gb18030")
            )
            convert_reader_assets.convert_file(
                {"extension": "html", "source_url": "https://huggingface.co/datasets/VoiceOfML/Test/resolve/rev/source.html"},
                source,
                target,
                work,
            )
            self.assertIn("列宁", target.read_text(encoding="utf-8"))
            self.assertIn('charset="utf-8"', target.read_text(encoding="utf-8").lower())

    def test_html_conversion_decodes_quoted_big5_charset(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source = work / "source.htm"
            source.write_bytes('<html><head><meta charset="Big5"></head><body>繁體中文</body></html>'.encode("big5"))
            output = convert_reader_assets.inline_html_resources(
                source,
                "https://huggingface.co/datasets/VoiceOfML/Test/resolve/rev/source.htm",
                work,
            )
            text = output.read_text(encoding="utf-8")
            self.assertIn("繁體中文", text)
            self.assertIn('charset="utf-8"', text.lower())
            self.assertNotIn("big5", text.lower())

    def test_undecorated_gb18030_prefers_chinese_over_cyrillic(self):
        with tempfile.TemporaryDirectory() as root:
            source = Path(root) / "source.html"
            source.write_bytes("<p>目录 人民日报 研究资料</p>".encode("gb18030"))
            self.assertIn("目录 人民日报 研究资料", convert_reader_assets.decode_html_source(source))

    def test_html_conversion_inlines_local_images_and_stylesheets_only(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source = work / "source.html"
            source.write_text(
                '<link rel="stylesheet" href="css/site.css">'
                '<img src="images/picture.png"><img src="https://evil.test/x.png">'
                '<img src="../secret.png">', encoding="utf-8",
            )
            resources = {
                "https://huggingface.co/datasets/VoiceOfML/Test/resolve/rev/css/site.css": b"body{color:red;background:url(x.png)}",
                "https://huggingface.co/datasets/VoiceOfML/Test/resolve/rev/images/picture.png": b"PNG",
            }

            def download(url, target, **_kwargs):
                data = resources[url]
                target.write_bytes(data)
                return hashlib.sha256(data).hexdigest(), len(data)

            with patch.object(convert_reader_assets, "download_source", side_effect=download):
                output = convert_reader_assets.inline_html_resources(
                    source,
                    "https://huggingface.co/datasets/VoiceOfML/Test/resolve/rev/source.html",
                    work,
                )
            text = output.read_text(encoding="utf-8")
            self.assertIn("data:image/png;base64", text)
            self.assertIn("<style>body{color:red;}</style>", text)
            self.assertNotIn('src="https://evil.test/x.png"', text)
            self.assertNotIn('src="../secret.png"', text)

    def test_page_content_ratio_detects_blank_and_nonblank_pages(self):
        from PIL import Image

        with tempfile.TemporaryDirectory() as root:
            blank = Path(root) / "blank.png"
            content = Path(root) / "content.png"
            Image.new("RGB", (100, 100), "white").save(blank)
            image = Image.new("RGB", (100, 100), "white")
            for x in range(20, 80):
                for y in range(20, 80):
                    image.putpixel((x, y), (0, 0, 0))
            image.save(content)
            self.assertEqual(convert_reader_assets.page_content_ratio(blank), 0)
            self.assertGreater(convert_reader_assets.page_content_ratio(content), 0.1)

    def test_pdf_content_validation_rejects_empty_pages_after_midpoint(self):
        from PIL import Image

        with tempfile.TemporaryDirectory() as root:
            work, pdf = Path(root), Path(root) / "document.pdf"
            pdf.write_bytes(b"%PDF-test")

            def render(command):
                page = int(command[command.index("-f") + 1])
                image = Image.new("RGB", (40, 40), "white")
                if page <= 2:
                    for x in range(10, 30):
                        for y in range(10, 30):
                            image.putpixel((x, y), (0, 0, 0))
                output = Path(command[-1]).with_suffix(".png")
                output.parent.mkdir(parents=True, exist_ok=True)
                image.save(output)

            with patch.object(convert_reader_assets, "command_output", return_value="Pages: 4\n"), \
                    patch.object(convert_reader_assets, "run_checked", side_effect=render):
                with self.assertRaisesRegex(RuntimeError, "after its midpoint"):
                    convert_reader_assets.validate_pdf_content(pdf, work)

    def test_endpoint_content_validation_rejects_empty_epub_docx_and_html(self):
        with tempfile.TemporaryDirectory() as root:
            root = Path(root)
            epub = root / "book.epub"
            with zipfile.ZipFile(epub, "w") as archive:
                archive.writestr("mimetype", "application/epub+zip")
                archive.writestr("META-INF/container.xml", "<container><rootfiles><rootfile full-path=\"content.opf\"/></rootfiles></container>")
                archive.writestr("content.opf", "<package><manifest><item id=\"chapter\" href=\"chapter.xhtml\"/></manifest><spine><itemref idref=\"chapter\"/></spine></package>")
                archive.writestr("chapter.xhtml", "<html><body><p></p></body></html>")
            with self.assertRaisesRegex(RuntimeError, "no readable content"):
                convert_reader_assets.validate_epub_content(epub)

            docx = root / "document.docx"
            with zipfile.ZipFile(docx, "w") as archive:
                archive.writestr("word/document.xml", "<w:document xmlns:w=\"urn:word\"><w:body><w:p/></w:body></w:document>")
            with self.assertRaisesRegex(RuntimeError, "no readable content"):
                convert_reader_assets.validate_docx_content(docx)

            html_file = root / "document.html"
            html_file.write_text("<html><body><style>only css</style></body></html>", encoding="utf-8")
            with self.assertRaisesRegex(RuntimeError, "no readable content"):
                convert_reader_assets.validate_html_content(html_file)

    def test_spreadsheet_page_manifest_uses_the_reader_compact_schema(self):
        with tempfile.TemporaryDirectory() as root:
            manifest = Path(root) / "page-manifest.json"
            manifest.write_text(json.dumps({
                "version": 2, "kind": "pdf-pages", "source_sha256": "a" * 64,
                "profile": reader_assets.SPREADSHEET_PAGE_PROFILE, "page_count": 3,
            }), encoding="utf-8")
            convert_reader_assets.validate_page_manifest(manifest)
            manifest.write_text(json.dumps({
                "version": 2, "kind": "pdf-pages", "page_count": 3, "pages": [],
            }), encoding="utf-8")
            with self.assertRaisesRegex(RuntimeError, "page manifest is invalid"):
                convert_reader_assets.validate_page_manifest(manifest)

    def test_xlsx_page_stream_exports_html_without_pdf_conversion(self):
        from openpyxl import Workbook
        from openpyxl.chart import LineChart, Reference
        from PIL import Image

        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source = work / "sample.xlsx"
            workbook = Workbook()
            first = workbook.active
            first.title = "Summary"
            first.append(["Year", "Value"])
            first.append([2024, 10])
            chart = LineChart()
            chart.add_data(Reference(first, min_col=2, min_row=1, max_row=2), titles_from_data=True)
            chart.set_categories(Reference(first, min_col=1, min_row=2, max_row=2))
            first.add_chart(chart, "D2")
            second = workbook.create_sheet("Details")
            second.append(["完整表格", "数据"])
            workbook.save(source)
            workbook.close()

            output_dir = work / "spreadsheet-html"
            def export_html(command, **_kwargs):
                html_output = Path(command[command.index("--outdir") + 1])
                html_output.mkdir(parents=True, exist_ok=True)
                (html_output / "sample.html").write_text(
                    '<a>Summary</a><a>Details</a><a href="sample_html_1.html">Summary</a>',
                    encoding="utf-8",
                )
                (html_output / "sample_html_1.html").write_text(
                    '<h1>Summary</h1><table><tr><td>Year</td><td>Value</td></tr>'
                    '<tr><td>2024</td><td>10</td></tr></table><img src="chart.png">',
                    encoding="utf-8",
                )
                (html_output / "sample_html_2.html").write_text(
                    '<h1>Details</h1><table><tr><td>完整表格</td><td>数据</td></tr></table>',
                    encoding="utf-8",
                )

            rendered = work / "rendered"
            rendered.mkdir()
            screenshot = rendered / "sheet.png"
            Image.new("RGB", (40, 40), "white").save(screenshot)
            target = work / "bundle/objects/aa/1234567890abcdef/page-manifest.json"
            item = {"profile": reader_assets.SPREADSHEET_PAGE_PROFILE}
            with patch.object(convert_reader_assets, "run_checked", side_effect=export_html) as run, \
                    patch.object(convert_reader_assets, "validate_office_pdf") as validate_pdf, \
                    patch.object(convert_reader_assets, "render_spreadsheet_html", return_value=[screenshot]):
                convert_reader_assets.convert_spreadsheet_to_pages(
                    source, target, work, item, "a" * 64, "objects/aa/" + "a" * 64
                    + "/1234567890abcdef/page-manifest.json",
                )

            command = run.call_args.args[0]
            self.assertIn("html:XHTML Calc File:UTF8", command)
            self.assertNotIn("pdf", command[command.index("--convert-to") + 1].lower())
            validate_pdf.assert_not_called()
            manifest = json.loads(target.read_text(encoding="utf-8"))
            self.assertEqual(manifest["page_count"], 1)
            self.assertNotIn("pages", manifest)

    def test_epub_content_validation_rejects_missing_image_resources(self):
        with tempfile.TemporaryDirectory() as root:
            epub = Path(root) / "missing-image.epub"
            with zipfile.ZipFile(epub, "w") as archive:
                archive.writestr("mimetype", "application/epub+zip")
                archive.writestr("META-INF/container.xml", '<container><rootfiles><rootfile full-path="O/content.opf"/></rootfiles></container>')
                archive.writestr("O/content.opf", '<package><manifest><item id="chapter" href="chapter.xhtml"/></manifest><spine><itemref idref="chapter"/></spine></package>')
                archive.writestr("O/chapter.xhtml", '<html><body><p>Readable text in this chapter.</p><img src="images/missing.png"/></body></html>')
            with self.assertRaisesRegex(RuntimeError, "image resource is missing"):
                convert_reader_assets.validate_epub_content(epub)

    def test_mhtml_content_validation_rejects_missing_or_invalid_images(self):
        with tempfile.TemporaryDirectory() as root:
            document = Path(root) / "document.html"
            document.write_text('<html><body><p>Readable saved page.</p><img src="image.png"></body></html>', encoding="utf-8")
            with self.assertRaisesRegex(RuntimeError, "unembedded image"):
                convert_reader_assets.validate_mhtml_content(document)

            document.write_text('<html><body><p>Readable saved page.</p><img alt="missing source"></body></html>', encoding="utf-8")
            with self.assertRaisesRegex(RuntimeError, "without a source"):
                convert_reader_assets.validate_mhtml_content(document)

            document.write_text('<html><body><p>Readable saved page.</p><img src="data:image/png;base64,aW1hZ2U="></body></html>', encoding="utf-8")
            with self.assertRaisesRegex(RuntimeError, "invalid embedded image"):
                convert_reader_assets.validate_mhtml_content(document)

            from PIL import Image
            image_bytes = io.BytesIO()
            Image.new("RGB", (2, 2), "white").save(image_bytes, format="PNG")
            encoded = base64.b64encode(image_bytes.getvalue()).decode("ascii")
            document.write_text(f'<html><body><p>Readable saved page.</p><img src="data:image/png;base64,{encoded}"></body></html>', encoding="utf-8")
            convert_reader_assets.validate_mhtml_content(document)

    def test_docx_validation_requires_document_structure(self):
        import zipfile

        with tempfile.TemporaryDirectory() as root:
            valid = Path(root) / "valid.docx"
            with zipfile.ZipFile(valid, "w") as archive:
                archive.writestr("[Content_Types].xml", "<Types/>")
                archive.writestr("word/document.xml", "<w:document/>")
            convert_reader_assets.validate_output(valid, "docx")
            invalid = Path(root) / "invalid.docx"
            with zipfile.ZipFile(invalid, "w") as archive:
                archive.writestr("other.xml", "<x/>")
            with self.assertRaises(RuntimeError):
                convert_reader_assets.validate_output(invalid, "docx")

    def test_docx_password_is_extracted_only_from_explicit_path_marker(self):
        self.assertEqual(reader_assets.source_password("repo", "资料〔密码：123〕.docx"), "123")
        self.assertEqual(reader_assets.source_password("repo", "book password: secret.docx"), "secret")
        self.assertEqual(reader_assets.source_password("repo", "ordinary.docx"), "")

    def test_password_protected_pdf_is_decrypted_with_qpdf(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.pdf", work / "document.pdf"
            source.write_bytes(b"%PDF-encrypted")
            with patch.object(convert_reader_assets, "run_checked") as run:
                convert_reader_assets.convert_file({
                    "extension": "pdf", "repo": "VoiceOfML/Test", "path": "book（密码123）.pdf",
                }, source, target, work)
            self.assertEqual(run.call_args.args[0][:3], ["qpdf", "--password=123", "--decrypt"])

    def test_ole_signature_is_detected_for_mislabeled_docx(self):
        self.assertEqual(convert_reader_assets.OLE_SIGNATURE, b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")

    def test_mislabeled_ole_docx_is_converted_with_doc_extension(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.docx", work / "document.docx"
            source.write_bytes(convert_reader_assets.OLE_SIGNATURE + b"data")
            def fake_run(command):
                (work / "mislabeled-office" / "mislabeled.docx").write_bytes(b"docx")
            with patch.object(convert_reader_assets, "run_checked", side_effect=fake_run) as run:
                convert_reader_assets.convert_file({"extension": "docx", "path": "book.docx"}, source, target, work)
            self.assertTrue(str(run.call_args.args[0][-1]).endswith("mislabeled.doc"))
            self.assertEqual(target.read_bytes(), b"docx")

    def test_conversion_command_timeout_is_bounded(self):
        with patch.object(convert_reader_assets.subprocess, "run", side_effect=convert_reader_assets.subprocess.TimeoutExpired("libreoffice", convert_reader_assets.COMMAND_TIMEOUT_SECONDS)):
            with self.assertRaisesRegex(RuntimeError, "timed out: libreoffice"):
                convert_reader_assets.run_checked(["libreoffice", "--headless"])

    def test_conversion_commands_do_not_inherit_publish_credentials(self):
        completed = Mock(returncode=0, stderr="")
        with patch.dict(convert_reader_assets.os.environ, {
                "HF_TOKEN": "hf-secret", "PAGES_TOKEN": "pages-secret",
                "READER_CONVERSION_PASSWORD": "office-secret", "SAFE_VALUE": "kept",
                }), patch.object(convert_reader_assets.subprocess, "run", return_value=completed) as run:
            convert_reader_assets.run_checked(["libreoffice", "--headless"])
        child_env = run.call_args.kwargs["env"]
        self.assertEqual(child_env["SAFE_VALUE"], "kept")
        for secret in ("HF_TOKEN", "PAGES_TOKEN", "READER_CONVERSION_PASSWORD"):
            self.assertNotIn(secret, child_env)

    def test_conversion_failures_use_stable_public_classes(self):
        self.assertEqual(convert_reader_assets.conversion_error_class(
            convert_reader_assets.ReaderConversionTimeout("timeout")), "conversion-timeout")
        self.assertEqual(convert_reader_assets.conversion_error_class(
            convert_reader_assets.ReaderConversionCommandError("failed")), "conversion-command-failed")
        self.assertEqual(convert_reader_assets.conversion_error_class(
            urllib.error.URLError("offline")), "source-download-failed")
        self.assertEqual(convert_reader_assets.conversion_error_class(
            TimeoutError("offline")), "source-download-failed")

    def test_conversion_workers_are_bounded(self):
        self.assertGreaterEqual(convert_reader_assets.CONVERSION_WORKERS, 1)

    def test_html_disguised_as_doc_uses_html_extension(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root)
            source, target = work / "source.doc", work / "document.docx"
            source.write_bytes(b"\xef\xbb\xbf<html><body>text</body></html>")
            def fake_run(command):
                if "odt" in command:
                    (work / "html-office" / "source.odt").write_bytes(b"odt")
                (work / "office" / "source.docx").write_bytes(b"docx")
            with patch.object(convert_reader_assets, "run_checked", side_effect=fake_run) as run:
                convert_reader_assets.convert_file({"extension": "doc"}, source, target, work)
            self.assertTrue(str(run.call_args_list[0].args[0][-1]).endswith("source.html"))
            self.assertTrue(str(run.call_args_list[1].args[0][-1]).endswith("source.odt"))

    def test_duplicate_source_digest_reuses_converted_artifact(self):
        item = {
            "key": "VoiceOfML/Test\0Book.docx", "extension": "docx",
            "source_url": "https://example.test/book.docx", "source_revision": "rev1",
            "profile": "libreoffice-pdf-v2", "reader_mode": "pdf", "output_name": "document.pdf",
        }
        with tempfile.TemporaryDirectory() as root:
            bundle = Path(root)

            def download(_url, target):
                target.write_bytes(b"same source")
                return "a" * 64, 11

            def convert(_item, _source, target, _work):
                target.write_bytes(b"%PDF-converted")

            with patch.object(convert_reader_assets, "download_source", side_effect=download), \
                    patch.object(convert_reader_assets, "convert_file", side_effect=convert) as conversion, \
                    patch.object(convert_reader_assets, "validate_office_pdf"), \
                    patch.object(convert_reader_assets, "validate_reader_content"):
                first = convert_reader_assets.convert_item(item, bundle)
                second = convert_reader_assets.convert_item({**item, "key": "VoiceOfML/Test\0Copy.docx"}, bundle)

            self.assertEqual(conversion.call_count, 1)
            self.assertEqual(first["path"], second["path"])
            self.assertIn("/libreoffice-pdf-v2/", first["path"])

    def test_identical_html_in_different_contexts_does_not_share_output(self):
        item = {
            "key": "VoiceOfML/Test\0A/Book.html", "extension": "html",
            "source_url": "https://example.test/A/Book.html", "source_revision": "rev1",
            "profile": "sanitized-html-v5", "reader_mode": "html", "output_name": "document.html",
        }
        with tempfile.TemporaryDirectory() as root:
            def download(_url, target):
                target.write_bytes(b"<p>same source document</p>")
                return "a" * 64, 27

            def convert(_item, _source, target, _work):
                target.write_text(_item["source_url"], encoding="utf-8")

            with patch.object(convert_reader_assets, "download_source", side_effect=download), \
                    patch.object(convert_reader_assets, "convert_file", side_effect=convert) as conversion, \
                    patch.object(convert_reader_assets, "validate_reader_content"):
                first = convert_reader_assets.convert_item(item, Path(root))
                second = convert_reader_assets.convert_item({
                    **item, "key": "VoiceOfML/Test\0B/Book.html",
                    "source_url": "https://example.test/B/Book.html",
                }, Path(root))
            self.assertEqual(conversion.call_count, 2)
            self.assertNotEqual(first["path"], second["path"])

    def test_concurrent_duplicate_source_digest_converts_once(self):
        item = {
            "key": "VoiceOfML/Test\0Book.djvu", "extension": "djvu",
            "source_url": "https://example.test/book.djvu", "source_revision": "rev1",
            "profile": "djvulibre-pdf-v1", "reader_mode": "pdf", "output_name": "document.pdf",
        }
        with tempfile.TemporaryDirectory() as root:
            bundle = Path(root)

            def download(_url, target):
                target.write_bytes(b"same source")
                return "a" * 64, 11

            def convert(_item, _source, target, _work):
                target.write_bytes(b"%PDF-converted")

            with patch.object(convert_reader_assets, "download_source", side_effect=download), \
                    patch.object(convert_reader_assets, "convert_file", side_effect=convert) as conversion, \
                    patch.object(convert_reader_assets, "validate_djvu_pdf"), \
                    patch.object(convert_reader_assets, "validate_reader_content"):
                with concurrent.futures.ThreadPoolExecutor(max_workers=2) as executor:
                    results = list(executor.map(
                        lambda queued: convert_reader_assets.convert_item(queued, bundle),
                        [item, {**item, "key": "VoiceOfML/Test\0Copy.djvu"}],
                    ))

            self.assertEqual(conversion.call_count, 1)
            self.assertEqual(results[0]["path"], results[1]["path"])
            self.assertEqual(results[0]["sha256"], results[1]["sha256"])

    def test_remote_artifact_is_reused_for_matching_source_and_profile(self):
        item = {
            "key": "VoiceOfML/Test\0Moved.djvu", "extension": "djvu",
            "source_url": "https://example.test/moved.djvu", "source_revision": "rev2",
            "profile": "djvulibre-pdf-v1", "reader_mode": "pdf", "output_name": "document.pdf",
        }
        digest = "a" * 64
        artifact = b"%PDF-reused"
        reusable = {f"{digest}\0djvulibre-pdf-v1": {
            "path": f"objects/aa/{digest}/djvulibre-pdf-v1/document.pdf",
            "bytes": len(artifact), "sha256": hashlib.sha256(artifact).hexdigest(),
        }}
        with tempfile.TemporaryDirectory() as root:
            def download(_url, target):
                target.write_bytes(b"source")
                return digest, 6

            def reuse(_url, target, _digest):
                target.write_bytes(artifact)

            with patch.object(convert_reader_assets, "download_source", side_effect=download), \
                    patch.object(convert_reader_assets, "download_existing", side_effect=reuse) as reused, \
                    patch.object(convert_reader_assets, "convert_file") as conversion, \
                    patch.object(convert_reader_assets, "validate_reader_content"):
                result = convert_reader_assets.convert_item(item, Path(root), reusable)

            conversion.assert_not_called()
            reused.assert_called_once()
            self.assertTrue(result["reused"])
            self.assertEqual(result["sha256"], hashlib.sha256(artifact).hexdigest())

    def test_pdf_staging_rehomes_a_reused_dataset_artifact(self):
        item = {
            "key": "VoiceOfML/Test\0Moved.djvu", "extension": "djvu",
            "source_url": "https://example.test/moved.djvu", "source_revision": "rev2",
            "profile": "djvulibre-pdf-v1", "reader_mode": "pdf", "output_name": "document.pdf",
            "bucket_staging": True,
        }
        digest = "a" * 64
        artifact = b"%PDF-reused"
        reusable = {f"{digest}\0djvulibre-pdf-v1": {
            "path": f"objects/aa/{digest}/djvulibre-pdf-v1/document.pdf",
            "bytes": len(artifact), "sha256": hashlib.sha256(artifact).hexdigest(),
        }}
        with tempfile.TemporaryDirectory() as root:
            def download(_url, target):
                target.write_bytes(b"source")
                return digest, 6

            with patch.object(convert_reader_assets, "download_source", side_effect=download), \
                    patch.object(convert_reader_assets, "download_existing", side_effect=lambda _url, target, _digest: target.write_bytes(artifact)), \
                    patch.object(convert_reader_assets, "validate_reader_content"):
                result = convert_reader_assets.convert_item(item, Path(root), reusable)

        self.assertTrue(result["path"].startswith("staging/pdf/"))
        self.assertTrue(result["reused"])

    def test_concurrent_remote_reuse_marks_both_results_reused(self):
        item = {
            "key": "VoiceOfML/Test\0One.djvu", "extension": "djvu",
            "source_url": "https://example.test/one.djvu", "source_revision": "rev2",
            "profile": "djvulibre-pdf-v1", "reader_mode": "pdf", "output_name": "document.pdf",
        }
        digest = "a" * 64
        artifact = b"%PDF-reused"
        reusable = {f"{digest}\0djvulibre-pdf-v1": {
            "path": "objects/existing.pdf", "bytes": len(artifact),
            "sha256": hashlib.sha256(artifact).hexdigest(), "reader_mode": "pdf",
        }}
        with tempfile.TemporaryDirectory() as root:
            def download(_url, target):
                target.write_bytes(b"source")
                return digest, 6

            def reuse(_url, target, _digest):
                target.write_bytes(artifact)

            with patch.object(convert_reader_assets, "download_source", side_effect=download), \
                    patch.object(convert_reader_assets, "download_existing", side_effect=reuse) as reused, \
                    patch.object(convert_reader_assets, "validate_reader_content"):
                with concurrent.futures.ThreadPoolExecutor(max_workers=2) as executor:
                    results = list(executor.map(
                        lambda queued: convert_reader_assets.convert_item(queued, Path(root), reusable),
                        [item, {**item, "key": "VoiceOfML/Test\0Two.djvu"}],
                    ))
        self.assertEqual(reused.call_count, 1)
        self.assertTrue(all(result["reused"] for result in results))

    def test_office_pdf_requires_embedded_fonts_and_extractable_cjk(self):
        item = {"path": "目录/中文.docx"}
        with patch.object(convert_reader_assets, "embedded_pdf_fonts", return_value=["NotoSansCJK"]), \
                patch.object(convert_reader_assets, "command_output", return_value="中文内容\n"):
            convert_reader_assets.validate_office_pdf(Path("document.pdf"), item, Path("."))

        with patch.object(convert_reader_assets, "embedded_pdf_fonts", return_value=[]), \
                patch.object(convert_reader_assets, "command_output", return_value="中文内容\n"), \
                patch.object(convert_reader_assets, "outline_pdf_fonts") as outline, \
                patch.object(convert_reader_assets, "validate_output"):
            convert_reader_assets.validate_office_pdf(Path("document.pdf"), item, Path("."))
        outline.assert_called_once()

        with patch.object(convert_reader_assets, "embedded_pdf_fonts", return_value=["NotoSansCJK"]), \
                patch.object(convert_reader_assets, "command_output", return_value="中文内容\n"):
            convert_reader_assets.validate_office_pdf(Path("document.pdf"), item, Path("."))

        with patch.object(convert_reader_assets, "embedded_pdf_fonts", return_value=["NotoSansCJK"]), \
                patch.object(convert_reader_assets, "command_output", return_value="□□□□\n"), \
                patch.object(convert_reader_assets, "rasterize_pdf") as rasterize, \
                patch.object(convert_reader_assets, "validate_output"):
            convert_reader_assets.validate_office_pdf(Path("document.pdf"), item, Path("."))
        rasterize.assert_called_once_with(Path("document.pdf"), Path("."))

    def test_office_pdf_rewrites_missing_embedded_fonts(self):
        with patch.object(convert_reader_assets, "embedded_pdf_fonts", return_value=[]), \
                patch.object(convert_reader_assets, "outline_pdf_fonts") as outline, \
                patch.object(convert_reader_assets, "validate_output"), \
                patch.object(convert_reader_assets, "command_output", return_value="中文内容\n"):
            convert_reader_assets.validate_office_pdf(
                Path("document.pdf"), {"path": "目录/中文.docx"}, Path("work")
            )
        outline.assert_called_once_with(Path("document.pdf"), Path("work"))

    def test_office_pdf_recovers_blank_pages_for_current_profiles(self):
        for profile in ("libreoffice-pdf-office-v2", "libreoffice-pdf-office-xlsx-v3"):
            with self.subTest(profile=profile), tempfile.TemporaryDirectory() as root:
                work = Path(root)
                source, target, candidate = work / "source.xlsx", work / "document.pdf", work / "candidate.pdf"
                source.write_bytes(b"office")
                target.write_bytes(b"bad")
                candidate.write_bytes(b"fixed")
                with patch.object(convert_reader_assets, "command_output", return_value=""), \
                        patch.object(convert_reader_assets, "rasterize_pdf", side_effect=[RuntimeError("blank interior page 2"), None]) as rasterize, \
                        patch.object(convert_reader_assets, "normalized_office_pdf", return_value=candidate) as normalize, \
                        patch.object(convert_reader_assets, "validate_output"):
                    convert_reader_assets.validate_office_pdf(
                        target, {"path": "中文.xlsx", "profile": profile}, work, source,
                    )
                normalize.assert_called_once_with(source, work)
                self.assertEqual(rasterize.call_count, 2)
                self.assertEqual(target.read_bytes(), b"fixed")

    def test_converter_returns_failure_when_the_entire_batch_fails(self):
        with tempfile.TemporaryDirectory() as root:
            root_path = Path(root)
            queue = root_path / "queue.json"
            bundle = root_path / "bundle"
            queue.write_text(json.dumps({"items": [{
                "key": "VoiceOfML/Test\0bad.docx", "repo": "VoiceOfML/Test", "path": "bad.docx",
                "extension": "docx", "source_revision": "rev", "profile": "docx-native-v2",
            }]}), encoding="utf-8")
            with patch.object(convert_reader_assets, "parse_args", return_value=SimpleNamespace(queue=queue, bundle=bundle, dry_run=False)), \
                    patch.object(convert_reader_assets, "convert_item", side_effect=RuntimeError("missing converter")):
                self.assertEqual(convert_reader_assets.main(), 1)
            result = json.loads((bundle / "bundle.json").read_text(encoding="utf-8"))["results"][0]
            self.assertEqual(result["status"], "failed")

class PublicationTests(unittest.TestCase):
    def setUp(self):
        patcher = patch.object(
            publish_reader_assets, "remote_pdf_manifest", return_value={"version": 1, "files": {}})
        patcher.start()
        self.addCleanup(patcher.stop)
        ocr_patcher = patch.object(
            publish_reader_assets, "remote_pdf_ocr_manifest", return_value={"version": 1, "files": {}})
        ocr_patcher.start()
        self.addCleanup(ocr_patcher.stop)
        range_patcher = patch.object(
            publish_reader_assets, "remote_state", return_value={"version": 1, "files": {}})
        range_patcher.start()
        self.addCleanup(range_patcher.stop)

    def make_bundle(self, root: str, result: dict) -> Path:
        bundle = Path(root)
        artifact = bundle / result["path"]
        artifact.parent.mkdir(parents=True)
        artifact.write_bytes(b"%PDF-converted")
        result["bytes"] = artifact.stat().st_size
        result["sha256"] = hashlib.sha256(artifact.read_bytes()).hexdigest()
        (bundle / "bundle.json").write_text(
            json.dumps({"version": 1, "results": [result]}), encoding="utf-8"
        )
        return bundle

    def test_publish_contains_artifact_manifest_and_sidecar(self):
        key = "VoiceOfML/Test\0A/Book.docx"
        result = {
            "key": key, "status": "ready", "source_revision": "rev1",
            "source_sha256": "a" * 64, "source_bytes": 10, "source_extension": "docx",
            "profile": "libreoffice-pdf-v2", "reader_mode": "pdf",
            "path": "objects/aa/source/document.pdf",
        }
        api = Mock()
        api.file_exists.return_value = False
        with tempfile.TemporaryDirectory() as root:
            bundle = self.make_bundle(root, result)
            manifest, operations = publish_reader_assets.build_publish(api, "vomebook/Test", bundle)
        self.assertEqual(manifest["files"][key]["status"], "ready")
        self.assertEqual(len(operations), 3)
        self.assertTrue(all(isinstance(operation, CommitOperationAdd) for operation in operations))
        self.assertEqual(
            {operation.path_in_repo for operation in operations},
            {result["path"], "manifest.json", "reader_assets.json.gz"},
        )

    def test_combine_bundles_keeps_results_and_artifact_roots(self):
        results = []
        with tempfile.TemporaryDirectory() as root:
            for index in range(2):
                result = {
                    "key": f"VoiceOfML/Test\\0Book-{index}.docx", "status": "ready",
                    "source_revision": "rev1", "source_sha256": f"{index + 1:064x}",
                    "source_bytes": 10, "source_extension": "docx", "profile": "docx-native-v2",
                    "reader_mode": "docx", "path": f"objects/{index:02x}/document.docx",
                }
                bundle = self.make_bundle(str(Path(root) / f"bundle-{index}"), result)
                results.append((bundle, result))
            data, roots = publish_reader_assets.combine_bundles([item[0] for item in results])
            merged = publish_reader_assets.materialize_bundles(
                [item[0] for item in results], data, Path(root) / "merged")
            merged_paths = [merged.joinpath(item[1]["path"]).is_file() for item in results]

        self.assertEqual([item["key"] for item in data["results"]],
                         [item[1]["key"] for item in results])
        self.assertEqual(roots[results[0][1]["path"]], results[0][0])
        self.assertEqual(roots[results[1][1]["path"]], results[1][0])
        self.assertEqual(merged_paths, [True, True])

    def test_static_and_pdf_document_outputs_are_bucket_upload_candidates(self):
        data = {"results": [
            {"status": "ready", "reader_mode": "docx", "path": "objects/aa/document.docx"},
            {"status": "ready", "reader_mode": "html", "path": "objects/bb/document.html"},
            {"status": "ready", "reader_mode": "foliate", "path": "objects/cc/document.epub"},
            {"status": "ready", "reader_mode": "audio", "path": "objects/dd/audio.mp3"},
        ]}
        self.assertEqual(publish_reader_assets.bucket_paths(data), [
            "objects/aa/document.docx", "objects/bb/document.html", "objects/cc/document.epub",
        ])

    def test_bucket_chapter_artifacts_and_manifest_use_ebook_prefix(self):
        chapter_manifest = "objects/aa/epub-chapters/chapter-manifest.json"
        data = {"version": 1, "bucket_migration": True, "results": [{
            "key": "VoiceOfML/Test\0book.epub", "status": "ready", "source_revision": "rev",
            "source_sha256": "a" * 64, "source_bytes": 1, "source_extension": "epub",
            "profile": "foliate-original-v1", "reader_mode": "foliate",
            "path": "objects/aa/document.epub", "bytes": 5,
            "sha256": hashlib.sha256(b"epub!").hexdigest(),
            "chapter_manifest": chapter_manifest,
        }]}
        with tempfile.TemporaryDirectory() as root:
            bundle = Path(root)
            (bundle / "objects/aa").mkdir(parents=True)
            (bundle / "objects/aa/document.epub").write_bytes(b"epub!")
            chapter_root = bundle / "objects/aa/epub-chapters"
            (chapter_root / "chapters").mkdir(parents=True)
            (chapter_root / "resources/chapter-0001").mkdir(parents=True)
            (chapter_root / "chapter-manifest.json").write_text("{}", encoding="utf-8")
            (chapter_root / "chapters/chapter-0001.xhtml").write_text("<html/>", encoding="utf-8")
            (chapter_root / "resources/chapter-0001/cover.png").write_bytes(b"png")
            artifacts = publish_reader_assets.artifact_files(data, {
                chapter_manifest: bundle, "objects/aa/document.epub": bundle,
            })
            self.assertIn("ebook-chapters/objects/aa/epub-chapters/chapter-manifest.json", artifacts)
            self.assertEqual(artifacts["ebook-chapters/objects/aa/epub-chapters/chapters/chapter-0001.xhtml"][1],
                             str(chapter_root / "chapters/chapter-0001.xhtml"))
            self.assertEqual(artifacts["ebook-chapters/objects/aa/epub-chapters/resources/chapter-0001/cover.png"][1],
                             str(chapter_root / "resources/chapter-0001/cover.png"))
            fallback_artifacts = publish_reader_assets.artifact_files(data, {}, bundle)
            self.assertIn("ebook-chapters/objects/aa/epub-chapters/resources/chapter-0001/cover.png", fallback_artifacts)
            api = Mock()
            api.file_exists.return_value = False
            manifest, _operations = publish_reader_assets.build_publish(
                api, "vomebook/Test", bundle, data_override=data,
                artifact_roots={chapter_manifest: bundle},
            )
        self.assertEqual(manifest["files"][data["results"][0]["key"]]["chapter_manifest"],
                         "ebook-chapters/objects/aa/epub-chapters/chapter-manifest.json")

    def test_pdf_to_dataset_upload_commits_the_pdf_to_reader_assets(self):
        result = {
            "key": "VoiceOfML/Test\0table.xlsx", "status": "ready",
            "reader_mode": "pdf", "pdf_to_dataset": True, "path": "objects/aa/document.pdf",
        }
        api = Mock()
        api.repo_info.return_value.sha = "parent"
        with tempfile.TemporaryDirectory() as root:
            bundle = Path(root)
            artifact = bundle / result["path"]
            artifact.parent.mkdir(parents=True)
            artifact.write_bytes(b"%PDF-table")
            publish_reader_assets.publish_dataset_pdf_artifacts(
                api, "vomebook/Reader-Assets",
                {"pdf_to_dataset": True, "results": [result]}, bundle,
            )
        api.create_commit.assert_called_once()
        operations = api.create_commit.call_args.kwargs["operations"]
        self.assertEqual([operation.path_in_repo for operation in operations], [result["path"]])

    def test_s3_reader_upload_uses_only_known_paths(self):
        fake_client = Mock()
        fake_client.head_object.return_value = {"ContentLength": 5}
        with tempfile.TemporaryDirectory() as root, patch.dict("os.environ", {
            "HF_S3_ACCESS_KEY_ID": "key", "HF_S3_SECRET_ACCESS_KEY": "secret",
            "HF_S3_NAMESPACE": "vomebook", "HF_S3_UPLOAD_WORKERS": "1",
        }), patch.object(publish_reader_assets, "_s3_client", return_value=fake_client):
            local = Path(root) / "document.docx"
            local.write_bytes(b"asset")
            publish_reader_assets.s3_upload_artifacts(
                {"objects/aa/document.docx": (Path(root), str(local))},
                "vomebook/pdf-pages",
            )
        fake_client.upload_file.assert_called_once()
        args, kwargs = fake_client.upload_file.call_args
        self.assertEqual(args[:3], (str(local), "pdf-pages", "objects/aa/document.docx"))
        self.assertEqual(kwargs["ExtraArgs"]["ContentType"],
                         "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        fake_client.head_object.assert_called_once_with(Bucket="pdf-pages", Key="objects/aa/document.docx")

    def test_bucket_migration_uploads_reused_dataset_objects_before_marking_bucket(self):
        key = "VoiceOfML/Test\0Saved.mht"
        result = {
            "key": key, "status": "ready", "source_revision": "rev1",
            "source_sha256": "a" * 64, "source_bytes": 10, "source_extension": "mht",
            "profile": "sanitized-mhtml-v6", "reader_mode": "html",
            "path": "objects/aa/document.html", "bytes": len(b"<p>Saved</p>"),
            "sha256": hashlib.sha256(b"<p>Saved</p>").hexdigest(),
        }
        existing = {"version": 1, "files": {key: {
            "status": "ready", "source_revision": "rev1", "source_sha256": "a" * 64,
            "source_extension": "mht", "profile": "sanitized-mhtml-v6", "reader_mode": "html",
            "path": "objects/aa/document.html", "bytes": len(b"<p>Saved</p>"),
            "sha256": hashlib.sha256(b"<p>Saved</p>").hexdigest(),
        }}}
        api = Mock()
        api.file_exists.return_value = False
        with tempfile.TemporaryDirectory() as root:
            bundle = Path(root)
            artifact = bundle / result["path"]
            artifact.parent.mkdir(parents=True)
            artifact.write_bytes(b"<p>Saved</p>")
            data = {"version": 1, "bucket_migration": True, "results": [result]}
            with patch.object(publish_reader_assets, "remote_manifest", return_value=existing):
                manifest, operations = publish_reader_assets.build_publish(
                    api, "vomebook/Reader-Assets", bundle, data_override=data,
                )
        entry = manifest["files"][key]
        self.assertEqual(entry["bucket"], reader_assets.READER_ASSETS_BUCKET)
        self.assertEqual({operation.path_in_repo for operation in operations}, {
            "objects/aa/document.html", "manifest.json", "reader_assets.json.gz",
        })

    def test_s3_reader_index_upload_accepts_string_temp_directory(self):
        fake_client = Mock()
        fake_client.head_object.return_value = {"ContentLength": 2}
        with tempfile.TemporaryDirectory() as root, patch.dict("os.environ", {
            "HF_S3_ACCESS_KEY_ID": "key", "HF_S3_SECRET_ACCESS_KEY": "secret",
            "HF_S3_NAMESPACE": "vomebook", "HF_S3_UPLOAD_WORKERS": "1",
        }), patch.object(publish_reader_assets, "_s3_client", return_value=fake_client):
            index = Path(root) / "reader-index" / "manifest.json"
            index.parent.mkdir()
            index.write_text("{}", encoding="utf-8")
            publish_reader_assets.s3_upload_tree(root, "vomebook/pdf-pages")
        fake_client.upload_file.assert_called_once()

    def test_failed_retry_does_not_replace_existing_ready_asset(self):
        key = "VoiceOfML/Test\0A/Book.docx"
        existing = {"status": "ready", "source_revision": "old", "profile": "libreoffice-pdf-v2",
                    "reader_mode": "pdf", "path": "objects/old/document.pdf"}
        api = Mock()
        api.file_exists.return_value = True
        with tempfile.TemporaryDirectory() as root:
            remote = Path(root) / "remote.json"
            remote.write_text(json.dumps({"version": 1, "files": {key: existing}}), encoding="utf-8")
            api.hf_hub_download.return_value = str(remote)
            bundle = Path(root) / "bundle"
            bundle.mkdir()
            (bundle / "bundle.json").write_text(json.dumps({"version": 1, "results": [{
                "key": key, "status": "failed", "source_revision": "new",
                "source_extension": "docx", "profile": "libreoffice-pdf-v2", "error": "RuntimeError",
            }]}), encoding="utf-8")
            manifest, operations = publish_reader_assets.build_publish(api, "vomebook/Test", bundle)
        self.assertEqual(
            {field: manifest["files"][key][field] for field in existing}, existing,
        )
        self.assertEqual(manifest["files"][key]["failed_source_revision"], "new")
        self.assertEqual(manifest["files"][key]["failed_profile"], "libreoffice-pdf-v2")
        self.assertEqual(len(operations), 2)

    def test_publish_removes_inactive_mapping_and_marks_orphan(self):
        removed_key = "VoiceOfML/Test\0Old/Book.djvu"
        existing = {
            "status": "ready", "source_revision": "rev1", "source_sha256": "a" * 64,
            "source_extension": "djvu", "profile": "djvulibre-pdf-v1", "reader_mode": "pdf",
            "path": "objects/aa/document.pdf", "bytes": 10, "sha256": "b" * 64,
        }
        api = Mock()
        api.file_exists.return_value = True
        with tempfile.TemporaryDirectory() as root:
            remote = Path(root) / "remote.json"
            remote.write_text(json.dumps({"version": 1, "files": {removed_key: existing}}), encoding="utf-8")
            api.hf_hub_download.return_value = str(remote)
            bundle = Path(root) / "bundle"
            bundle.mkdir()
            (bundle / "bundle.json").write_text(json.dumps({
            "version": 1, "results": [], "active_keys": ["VoiceOfML/Test\0Other.djvu"],
            "authoritative_snapshot": True,
            }), encoding="utf-8")
            manifest, operations = publish_reader_assets.build_publish(api, "vomebook/Test", bundle)

        self.assertNotIn(removed_key, manifest["files"])
        self.assertEqual(manifest["orphans"][existing["path"]]["since"], date.today().isoformat())
        self.assertEqual(len(operations), 2)

    def test_empty_active_key_snapshot_removes_all_mappings(self):
        key = "VoiceOfML/Test\0Deleted.djvu"
        existing = {"status": "ready", "path": "objects/deleted", "source_sha256": "a" * 64}
        api = Mock()
        api.file_exists.return_value = True
        with tempfile.TemporaryDirectory() as root:
            remote = Path(root) / "remote.json"
            remote.write_text(json.dumps({"version": 1, "files": {key: existing}}), encoding="utf-8")
            api.hf_hub_download.return_value = str(remote)
            bundle = Path(root) / "bundle"
            bundle.mkdir()
            (bundle / "bundle.json").write_text(json.dumps({
                "version": 1, "results": [], "active_keys": [], "authoritative_snapshot": True,
            }), encoding="utf-8")
            manifest, _ = publish_reader_assets.build_publish(api, "vomebook/Test", bundle)
        self.assertEqual(manifest["files"], {})
        self.assertIn(existing["path"], manifest["orphans"])
    def test_reused_result_restores_orphan_without_upload(self):
        path = "objects/aa/document.pdf"
        result = {
            "key": "VoiceOfML/Test\0Moved.djvu", "status": "ready", "source_revision": "rev2",
            "source_sha256": "a" * 64, "source_bytes": 10, "source_extension": "djvu",
            "profile": "djvulibre-pdf-v1", "reader_mode": "pdf", "path": path,
            "bytes": len(b"%PDF-reused"), "sha256": hashlib.sha256(b"%PDF-reused").hexdigest(),
            "reused": True,
        }
        api = Mock()
        api.file_exists.return_value = True
        with tempfile.TemporaryDirectory() as root:
            remote = Path(root) / "remote.json"
            remote.write_text(json.dumps({"version": 1, "files": {}, "orphans": {
                path: {**result, "since": "2026-01-01"},
            }}), encoding="utf-8")
            api.hf_hub_download.return_value = str(remote)
            bundle = Path(root) / "bundle"
            artifact = bundle / path
            artifact.parent.mkdir(parents=True)
            artifact.write_bytes(b"%PDF-reused")
            (bundle / "bundle.json").write_text(json.dumps({
                "version": 1, "results": [result], "active_keys": [result["key"]],
                "authoritative_snapshot": True,
            }), encoding="utf-8")
            manifest, operations = publish_reader_assets.build_publish(api, "vomebook/Test", bundle)

        self.assertNotIn(path, manifest["orphans"])
        self.assertEqual(manifest["files"][result["key"]]["path"], path)
        self.assertEqual(len(operations), 2)

    def test_stale_reused_hint_uploads_when_object_is_absent_from_latest_manifest(self):
        result = {
            "key": "VoiceOfML/Test\0Moved.docx", "status": "ready", "source_revision": "rev2",
            "source_sha256": "a" * 64, "source_bytes": 10, "source_extension": "docx",
            "profile": "docx-native-v2", "reader_mode": "docx",
            "path": "objects/aa/document.docx", "reused": True,
        }
        api = Mock()
        api.file_exists.return_value = False
        with tempfile.TemporaryDirectory() as root:
            bundle = self.make_bundle(root, result)
            _, operations = publish_reader_assets.build_publish(api, "vomebook/Test", bundle)
        self.assertEqual(len(operations), 3)
        self.assertIn(result["path"], {operation.path_in_repo for operation in operations})

    def test_publish_reuses_matching_object_added_by_another_shard(self):
        path = "objects/aa/remote/document.pdf"
        remote_entry = {
            "status": "ready", "source_revision": "other", "source_sha256": "a" * 64,
            "source_bytes": 10, "source_extension": "docx", "profile": "libreoffice-pdf-v2",
            "reader_mode": "pdf", "path": path, "bytes": 12, "sha256": "b" * 64,
        }
        result = {
            **remote_entry, "key": "VoiceOfML/Test\0Copy.docx", "source_revision": "rev2",
            "path": "objects/aa/local/document.pdf", "bytes": 99, "sha256": "c" * 64,
        }
        api = Mock()
        api.file_exists.return_value = True
        with tempfile.TemporaryDirectory() as root:
            remote = Path(root) / "remote.json"
            remote.write_text(json.dumps({
                "version": 1, "files": {"VoiceOfML/Test\0Original.docx": remote_entry},
            }), encoding="utf-8")
            api.hf_hub_download.return_value = str(remote)
            bundle = Path(root) / "bundle"
            bundle.mkdir()
            (bundle / "bundle.json").write_text(json.dumps({"version": 1, "results": [result]}), encoding="utf-8")
            manifest, operations = publish_reader_assets.build_publish(api, "vomebook/Test", bundle)
        published = manifest["files"][result["key"]]
        self.assertEqual(published["path"], path)
        self.assertEqual(published["sha256"], remote_entry["sha256"])
        self.assertEqual(len(operations), 2)

    def test_force_publish_does_not_reuse_matching_remote_object(self):
        key = "VoiceOfML/Test\0Book.docx"
        remote_entry = {
            "status": "ready", "source_revision": "rev1", "source_sha256": "a" * 64,
            "source_bytes": 10, "source_extension": "docx", "profile": "libreoffice-pdf-v2",
            "reader_mode": "pdf", "path": "objects/aa/old/document.pdf", "bytes": 12, "sha256": "b" * 64,
        }
        result = {**remote_entry, "key": key, "path": "objects/aa/new/document.pdf"}
        api = Mock()
        api.file_exists.return_value = True
        with tempfile.TemporaryDirectory() as root:
            remote = Path(root) / "remote.json"
            remote.write_text(json.dumps({"version": 1, "files": {key: remote_entry}}), encoding="utf-8")
            api.hf_hub_download.return_value = str(remote)
            bundle = self.make_bundle(str(Path(root) / "bundle"), result)
            bundle_data = json.loads((bundle / "bundle.json").read_text(encoding="utf-8"))
            bundle_data["force_rebuild"] = True
            (bundle / "bundle.json").write_text(json.dumps(bundle_data), encoding="utf-8")
            manifest, operations = publish_reader_assets.build_publish(api, "vomebook/Test", bundle)
        self.assertEqual(manifest["files"][key]["path"], result["path"])
        self.assertEqual(len(operations), 3)

    def test_force_publish_refreshes_every_mapping_for_an_overwritten_object(self):
        path = "objects/aa/shared/document.pdf"
        old = {
            "status": "ready", "source_revision": "rev1", "source_sha256": "a" * 64,
            "source_bytes": 10, "source_extension": "docx", "profile": "libreoffice-pdf-v2",
            "reader_mode": "pdf", "path": path, "bytes": 12, "sha256": "b" * 64,
        }
        rebuilt_key = "VoiceOfML/Test\0Rebuilt.docx"
        shared_key = "VoiceOfML/Test\0Shared.docx"
        result = {**old, "key": rebuilt_key}
        api = Mock()
        api.file_exists.return_value = True
        with tempfile.TemporaryDirectory() as root:
            remote = Path(root) / "remote.json"
            remote.write_text(json.dumps({
                "version": 1, "files": {rebuilt_key: old, shared_key: {**old, "source_revision": "rev2"}},
            }), encoding="utf-8")
            api.hf_hub_download.return_value = str(remote)
            bundle = self.make_bundle(str(Path(root) / "bundle"), result)
            bundle_data = json.loads((bundle / "bundle.json").read_text(encoding="utf-8"))
            bundle_data["force_rebuild"] = True
            (bundle / "bundle.json").write_text(json.dumps(bundle_data), encoding="utf-8")
            manifest, _ = publish_reader_assets.build_publish(api, "vomebook/Test", bundle)
        self.assertEqual(manifest["files"][shared_key]["bytes"], result["bytes"])
        self.assertEqual(manifest["files"][shared_key]["sha256"], result["sha256"])

    def test_parent_409_conflict_rebuilds_publish_against_latest_revision(self):
        result = {
            "key": "VoiceOfML/Test\0A/Book.docx", "status": "ready", "source_revision": "rev1",
            "source_sha256": "a" * 64, "source_bytes": 10, "source_extension": "docx",
            "profile": "libreoffice-pdf-v2", "reader_mode": "pdf", "path": "objects/aa/document.pdf",
        }
        response = requests.Response()
        response.status_code = 409
        response.request = requests.Request("POST", "https://huggingface.co/api/datasets/vomebook/Test/commit/main").prepare()
        conflict = HfHubHTTPError("conflict", response=response)
        api = Mock()
        api.repo_info.side_effect = [Mock(sha="parent-1"), Mock(sha="uploaded-1"), Mock(sha="parent-2")]
        api.file_exists.return_value = False
        api.create_commit.side_effect = [conflict, None]
        with tempfile.TemporaryDirectory() as root:
            bundle = self.make_bundle(root, result)
            with patch.object(publish_reader_assets.time, "sleep"):
                _, count = publish_reader_assets.publish_bundle(api, "vomebook/Test", bundle)
        self.assertEqual(count, 1)
        self.assertEqual(api.create_commit.call_count, 2)
        self.assertEqual(
            [call.kwargs["parent_commit"] for call in api.create_commit.call_args_list],
            ["uploaded-1", "parent-2"],
        )

    def test_transient_upload_error_rebuilds_commit_with_fresh_signature(self):
        result = {
            "key": "VoiceOfML/Test\0A/Book.docx", "status": "ready", "source_revision": "rev1",
            "source_sha256": "a" * 64, "source_bytes": 10, "source_extension": "docx",
            "profile": "libreoffice-pdf-v2", "reader_mode": "pdf", "path": "objects/aa/document.pdf",
        }
        response = requests.Response()
        response.status_code = 503
        response.request = requests.Request("POST", "https://huggingface.co/api/datasets/vomebook/Test/commit/main").prepare()
        api = Mock()
        api.repo_info.side_effect = [Mock(sha="parent-1"), Mock(sha="uploaded-1"), Mock(sha="parent-2")]
        api.file_exists.return_value = False
        api.create_commit.side_effect = [HfHubHTTPError("unavailable", response=response), None]
        with tempfile.TemporaryDirectory() as root:
            bundle = self.make_bundle(root, result)
            with patch.object(publish_reader_assets.time, "sleep"):
                _, count = publish_reader_assets.publish_bundle(api, "vomebook/Test", bundle)
        self.assertEqual(count, 1)
        self.assertEqual(api.create_commit.call_count, 2)
        self.assertEqual(
            [call.kwargs["parent_commit"] for call in api.create_commit.call_args_list],
            ["uploaded-1", "parent-2"],
        )

    def test_transient_preflight_error_is_retried(self):
        result = {
            "key": "VoiceOfML/Test\0A/Book.docx", "status": "ready", "source_revision": "rev1",
            "source_sha256": "a" * 64, "source_bytes": 10, "source_extension": "docx",
            "profile": "docx-native-v2", "reader_mode": "docx", "path": "objects/aa/document.docx",
            "bytes": 12, "sha256": "b" * 64,
        }
        response = requests.Response()
        response.status_code = 429
        response.request = requests.Request("GET", "https://huggingface.co/api/datasets/vomebook/Test").prepare()
        api = Mock()
        api.repo_info.side_effect = [HfHubHTTPError("rate limited", response=response), Mock(sha="parent-1"), Mock(sha="uploaded-1")]
        api.file_exists.return_value = True
        with tempfile.TemporaryDirectory() as root, patch.object(publish_reader_assets.time, "sleep"):
            remote = Path(root) / "remote.json"
            remote.write_text(json.dumps({"version": 1, "files": {}}), encoding="utf-8")
            api.hf_hub_download.return_value = str(remote)
            bundle = self.make_bundle(root, result)
            _, count = publish_reader_assets.publish_bundle(api, "vomebook/Test", bundle)
        self.assertEqual(count, 1)
        self.assertEqual(api.repo_info.call_count, 3)

    def test_transient_error_after_remote_commit_is_idempotent(self):
        result = {
            "key": "VoiceOfML/Test\0A/Book.docx", "status": "ready", "source_revision": "rev1",
            "source_sha256": "a" * 64, "source_bytes": 10, "source_extension": "docx",
            "profile": "libreoffice-pdf-v2", "reader_mode": "pdf", "path": "objects/aa/document.pdf",
            "bytes": 12, "sha256": "b" * 64,
        }
        response = requests.Response()
        response.status_code = 504
        response.request = requests.Request("POST", "https://huggingface.co/api/datasets/vomebook/Test/commit/main").prepare()
        api = Mock()
        api.repo_info.side_effect = [Mock(sha="parent-1"), Mock(sha="uploaded-1"), Mock(sha="parent-2")]
        api.file_exists.return_value = True
        with tempfile.TemporaryDirectory() as root:
            bundle = self.make_bundle(root, result)
            remote = Path(root) / "remote.json"
            remote.write_text(json.dumps({"version": 1, "files": {result["key"]: result}}), encoding="utf-8")
            api.hf_hub_download.return_value = str(remote)
            api.create_commit.side_effect = HfHubHTTPError("gateway timeout", response=response)
            with patch.object(publish_reader_assets.time, "sleep"):
                _, count = publish_reader_assets.publish_bundle(api, "vomebook/Test", bundle)
        self.assertEqual(count, 0)
        self.assertEqual(api.create_commit.call_count, 1)

    def test_parent_retry_rejects_a_concurrent_change_to_the_same_key(self):
        key = "VoiceOfML/Test\0A/Book.docx"
        result = {
            "key": key, "status": "ready", "source_revision": "rev1", "source_sha256": "a" * 64,
            "source_bytes": 10, "source_extension": "docx", "profile": "docx-native-v2",
            "reader_mode": "docx", "path": "objects/aa/document.docx",
        }
        response = requests.Response()
        response.status_code = 412
        response.request = requests.Request("POST", "https://huggingface.co/api/datasets/vomebook/Test/commit/main").prepare()
        api = Mock()
        api.repo_info.side_effect = [Mock(sha="parent-1"), Mock(sha="uploaded-1"), Mock(sha="parent-2")]
        api.file_exists.return_value = True
        with tempfile.TemporaryDirectory() as root:
            first = Path(root) / "first.json"
            second = Path(root) / "second.json"
            first.write_text(json.dumps({"version": 1, "files": {}}), encoding="utf-8")
            second.write_text(json.dumps({"version": 1, "files": {key: {
                **result, "source_revision": "newer", "path": "objects/bb/newer.docx",
            }}}), encoding="utf-8")
            api.hf_hub_download.side_effect = [str(first), str(first), str(second)]
            api.create_commit.side_effect = HfHubHTTPError("conflict", response=response)
            bundle = self.make_bundle(str(Path(root) / "bundle"), result)
            with patch.object(publish_reader_assets.time, "sleep"), self.assertRaisesRegex(
                    RuntimeError, "key changed"):
                publish_reader_assets.publish_bundle(api, "vomebook/Test", bundle)

    def test_legacy_bundle_without_authoritative_marker_never_removes_mappings(self):
        key = "VoiceOfML/Test\0Keep.djvu"
        existing = {"status": "ready", "path": "objects/keep", "reader_mode": "pdf"}
        api = Mock()
        api.file_exists.return_value = True
        with tempfile.TemporaryDirectory() as root:
            remote = Path(root) / "remote.json"
            remote.write_text(json.dumps({"version": 1, "files": {key: existing}}), encoding="utf-8")
            api.hf_hub_download.return_value = str(remote)
            bundle = Path(root) / "bundle"
            bundle.mkdir()
            (bundle / "bundle.json").write_text(json.dumps({
                "version": 1, "results": [], "active_keys": [],
            }), encoding="utf-8")
            manifest, _ = publish_reader_assets.build_publish(api, "vomebook/Test", bundle)
        self.assertEqual(manifest["files"][key], existing)

    def test_replacing_ready_mapping_marks_old_object_orphan(self):
        key = "VoiceOfML/Test\0Changed.djvu"
        previous = {"status": "ready", "path": "objects/old", "source_sha256": "a" * 64}
        result = {
            "key": key, "status": "ready", "source_revision": "rev2", "source_sha256": "b" * 64,
            "source_bytes": 12, "source_extension": "djvu", "profile": "djvulibre-pdf-v1",
            "reader_mode": "pdf", "path": "objects/new",
        }
        api = Mock()
        api.file_exists.return_value = True
        with tempfile.TemporaryDirectory() as root:
            remote = Path(root) / "remote.json"
            remote.write_text(json.dumps({"version": 1, "files": {key: previous}}), encoding="utf-8")
            api.hf_hub_download.return_value = str(remote)
            bundle = self.make_bundle(root + "/bundle", result)
            manifest, _ = publish_reader_assets.build_publish(api, "vomebook/Test", bundle)
        self.assertIn(previous["path"], manifest["orphans"])

    def test_sidecar_is_compact_and_deterministic(self):
        manifest = {"version": 1, "files": {
            "b": {"status": "failed"},
            "a": {"status": "ready", "reader_mode": "epub", "path": "objects/a/book.epub"},
        }}
        one = build_reader_assets_index.encode_index(manifest)
        two = build_reader_assets_index.encode_index(manifest)
        self.assertEqual(one, two)
        self.assertEqual(json.loads(gzip.decompress(one)), {
            "v": 1, "f": {"a": {"s": 2, "m": "e", "p": "objects/a/book.epub"}, "b": {"s": 4}}
        })

    def test_sidecar_encodes_docx_reader_mode(self):
        manifest = {"version": 1, "files": {"book": {
            "status": "ready", "reader_mode": "docx", "path": "objects/aa/source/docx-native-v1/document.docx",
            "bucket": reader_assets.READER_ASSETS_BUCKET,
        }}}
        self.assertEqual(build_reader_assets_index.build_index(manifest)["f"]["book"], {
            "s": 2, "m": "d", "p": "objects/aa/source/docx-native-v1/document.docx",
            "b": reader_assets.READER_ASSETS_BUCKET,
        })

    def test_sidecar_encodes_html_reader_mode(self):
        manifest = {"version": 1, "files": {"page": {
            "status": "ready", "reader_mode": "html", "path": "objects/aa/source/sanitized-html-v1/document.html",
        }}}
        self.assertEqual(build_reader_assets_index.build_index(manifest)["f"]["page"], {
            "s": 2, "m": "h", "p": "objects/aa/source/sanitized-html-v1/document.html",
        })

    def test_sidecar_encodes_audio_and_video_reader_modes(self):
        manifest = {"version": 1, "files": {
            "sound": {"status": "ready", "reader_mode": "audio", "path": "objects/aa/source/audio.mp3"},
            "movie": {"status": "ready", "reader_mode": "video", "path": "objects/bb/source/video.mp4"},
        }}
        files = build_reader_assets_index.build_index(manifest)["f"]
        self.assertEqual(files["sound"]["m"], "a")
        self.assertEqual(files["movie"]["m"], "v")

    def test_sidecar_merges_ready_pdf_assets(self):
        manifest = {"version": 1, "files": {
            "book": {"status": "ready", "reader_mode": "epub", "path": "objects/aa/book.epub"},
        }}
        pdf_manifest = {"version": 1, "files": {
            "scan": {"status": "ready", "strategy": "sampled-webp",
                     "render_profile": pdf_assets.PDF_PROFILE,
                     "decision_profile": pdf_assets.PDF_DECISION_PROFILE,
                     "path": "", "page_manifest": {
                "path": "objects/bb/" + "b" * 64 + "/page-manifest.json",
            }},
            "text": {"status": "skipped", "reason": "native-text-pdf"},
        }}
        files = build_reader_assets_index.build_index(manifest, pdf_manifest)["f"]
        self.assertEqual(files["book"]["m"], "e")
        self.assertEqual(files["scan"], {
            "s": 2, "m": "p", "p": "objects/bb/" + "b" * 64 + "/page-manifest.json",
            "b": "vomebook/pdf-pages",
        })
        self.assertNotIn("text", files)

    def test_sidecar_preserves_pdf_ocr_when_reader_assets_change(self):
        manifest = {"version": 1, "files": {
            "scan": {"status": "ready", "reader_mode": "pdf",
                      "path": "objects/aa/source/document.pdf"},
        }}
        ocr = {"version": 1, "files": {
            "scan": {"status": "ready", "profile": "pdf-ocr-profile",
                      "classification": "scan",
                      "ocr_manifest": "objects/aa/" + "a" * 64 + "/" + "b" * 16 + "/ocr-manifest.json"},
        }}
        self.assertEqual(build_reader_assets_index.build_index(manifest, ocr_manifest=ocr)["f"]["scan"], {
            "s": 2, "m": "p", "p": "objects/aa/source/document.pdf",
            "o": ocr["files"]["scan"]["ocr_manifest"], "om": "scan",
            "ob": "vomebook/pdf-pages",
        })

    def test_sidecar_excludes_legacy_streams_and_linearized_pdfs(self):
        manifest = {"version": 1, "files": {}}
        pdf_manifest = {"version": 1, "files": {
            "legacy": {"status": "ready", "strategy": "sampled-webp",
                       "page_manifest": {"path": "objects/aa/" + "a" * 64 + "/page-manifest.json"}},
            "current": {"status": "ready", "strategy": "sampled-webp", "source_extension": "caj",
                        "render_profile": pdf_assets.PDF_PROFILE,
                        "decision_profile": pdf_assets.PDF_DECISION_PROFILE,
                        "page_manifest": {"path": "objects/bb/" + "b" * 64 + "/page-manifest.json"}},
            "linear": {"status": "ready", "strategy": "linearized-pdf", "source_extension": "pdf",
                       "path": "objects/cc/" + "c" * 64 + "/linearized.pdf"},
        }}
        files = build_reader_assets_index.build_index(manifest, pdf_manifest)["f"]
        self.assertNotIn("legacy", files)
        self.assertIn("current", files)
        self.assertNotIn("linear", files)

    def test_reader_workflow_uses_explicit_empty_queue_guard_under_errexit(self):
        workflow = (Path(__file__).parents[1] / ".github/workflows/reader-assets.yml").read_text(encoding="utf-8")
        self.assertIn('if [[ "${count}" == "0" ]]; then\n            exit 0\n          fi', workflow)
        self.assertNotIn('[[ "${count}" == "0" ]] && exit 0', workflow)


class SearchIndexPublicationTests(unittest.TestCase):
    def test_pages_publish_copies_only_sidecar_and_pushes(self):
        commands = []

        def fake_run(args, cwd=None, env=None):
            commands.append((list(args), cwd, env))
            if args[:2] == ["git", "clone"]:
                Path(args[-1], "data").mkdir(parents=True, exist_ok=True)
                return 0, "", ""
            if args[:3] == ["git", "status", "--porcelain"]:
                return 0, " M data/reader_assets.json.gz", ""
            return 0, "", ""

        with tempfile.TemporaryDirectory() as root:
            source = Path(root) / "reader_assets.json.gz"
            source.write_bytes(b"gzip")
            with patch.object(publish_search_reader_index, "run", side_effect=fake_run):
                publish_search_reader_index.publish_to_pages(
                    source, "vomebook/search", "token", dry_run=False
                )
        argv = [item[0] for item in commands]
        self.assertIn(["git", "add", "data/reader_assets.json.gz"], argv)
        commit = next(args for args in argv if args[:2] == ["git", "commit"])
        self.assertNotIn("skip ci", " ".join(commit).lower())
        self.assertEqual(argv[-1], ["git", "push"])

    def test_dry_run_never_uses_remote_publishers(self):
        publish_search_reader_index.publish_to_pages(
            Path("missing"), "", "", dry_run=True
        )


class PruneTests(unittest.TestCase):
    def test_only_unreferenced_orphans_past_grace_period_expire(self):
        manifest = {"version": 1, "files": {
            "live": {"status": "ready", "path": "objects/live"},
        }, "orphans": {
            "objects/old": {"since": "2026-06-01"},
            "objects/new": {"since": "2026-08-20"},
            "objects/live": {"since": "2026-01-01"},
            "objects/invalid": {"since": "unknown"},
        }}
        self.assertEqual(
            prune_reader_assets.expired_orphans(manifest, date(2026, 8, 26), 30, 100),
            ["objects/old"],
        )

    def test_zero_day_grace_deletes_today_orphans_but_not_referenced_objects(self):
        manifest = {"version": 1, "files": {
            "live": {"status": "ready", "path": "objects/live"},
        }, "orphans": {
            "objects/today": {"since": "2026-08-26"},
            "objects/live": {"since": "2026-01-01"},
            "objects/invalid": {"since": "unknown"},
        }}
        self.assertEqual(
            prune_reader_assets.expired_orphans(manifest, date(2026, 8, 26), 0, 100),
            ["objects/today"],
        )

    def test_prune_cli_accepts_zero_day_grace(self):
        with patch.object(sys, "argv", ["prune_reader_assets.py", "--grace-days", "0"]):
            args = prune_reader_assets.parse_args()
        self.assertEqual(args.grace_days, 0)

    def test_prune_deletes_objects_and_republishes_manifest_and_sidecar(self):
        manifest = {"version": 1, "files": {}, "orphans": {
            "objects/old": {"since": "2026-01-01"},
            "objects/keep": {"since": "2026-08-20"},
        }}
        updated, operations = prune_reader_assets.build_prune(manifest, ["objects/old"])
        self.assertNotIn("objects/old", updated["orphans"])
        self.assertIn("objects/keep", updated["orphans"])
        self.assertEqual(
            {operation.path_in_repo for operation in operations},
            {"objects/old", "manifest.json", "reader_assets.json.gz"},
        )
        self.assertEqual(sum(isinstance(operation, CommitOperationDelete) for operation in operations), 1)


class WorkflowContractTests(unittest.TestCase):
    def test_workflow_exposes_incremental_controls_and_excludes_pdg(self):
        workflow = Path(".github/workflows/reader-assets.yml").read_text(encoding="utf-8")
        for field in ("repo:", "extension:", "limit:", "checkpoint_batches:", "retry_failed:", "force:", "dry_run:"):
            self.assertIn(field, workflow)
        self.assertNotIn("pdg", workflow.lower())
        self.assertIn("publish_search_reader_index.py", workflow)
        self.assertIn("fonts-noto-cjk", workflow)
        self.assertIn("fonts-wqy-microhei", workflow)
        self.assertIn("ghostscript", workflow)
        self.assertIn("poppler-utils", workflow)
        self.assertIn("tesseract-ocr-chi-sim", workflow)
        self.assertIn("djvulibre-bin", workflow)
        self.assertIn("doc, docx, epub, htm, html, mobi, azw3, fb2, odt, rtf", workflow)
        self.assertIn("chm, tif, tiff, djvu, ppt, pptx, pps, odp", workflow)
        self.assertIn("htm|html) packages=()", workflow)
        self.assertIn("inputs.limit || '20'", workflow)
        self.assertIn("inputs.checkpoint_batches || '30'", workflow)
        self.assertIn("python scripts/publish_reader_assets.py", workflow)
        self.assertIn("python scripts/prepare_reader_assets_plan.py", workflow)
        self.assertNotIn("python - <<'PY'", workflow)
        self.assertIn("djvu) packages=(djvulibre-bin qpdf poppler-utils)", workflow)
        self.assertIn("epub) packages=(calibre qpdf)", workflow)
        self.assertIn("mobi|azw3|fb2) packages=(calibre)", workflow)
        self.assertIn('python scripts/publish_reader_assets.py --bundles "${bundles[@]}"', workflow)
        self.assertIn("odt) packages=(libreoffice)", workflow)
        self.assertIn("rtf) packages=(calibre libreoffice", workflow)
        self.assertIn("chm) packages=(calibre p7zip-full)", workflow)
        self.assertIn("tif|tiff) packages=(poppler-utils)", workflow)
        self.assertIn("mht|mhtml) packages=()", workflow)
        self.assertIn("pip install -r scripts/requirements-spreadsheet-render.txt", workflow)
        self.assertIn("playwright install --with-deps chromium", workflow)
        staging_section = workflow.split("      bucket_pdf_staging:\n", 1)[1].split("      dry_run:\n", 1)[0]
        self.assertIn("default: false", staging_section)
        self.assertIn("pdf-optimized", staging_section)
        self.assertIn("ps) packages=(ghostscript poppler-utils)", workflow)
        self.assertIn("caj|kdh) packages=(git mupdf-tools poppler-utils", workflow)
        self.assertIn("checkout --detach 6c4bc32b15ce748d211f45d536f5d5511ef9f368", workflow)
        self.assertIn("CAJ2PDF_DIR: /opt/caj2pdf", workflow)
        self.assertIn("ape|wma|amr|flv|f4v|rm|rmvb|mkv|avi|mpg|mpeg|mts|ts|wmv) packages=(ffmpeg)", workflow)
        self.assertIn("READER_CONVERSION_WORKERS:", workflow)
        self.assertIn("needs.plan.outputs.extension == 'djvu'", workflow)
        self.assertIn("needs.plan.outputs.count != '0'", workflow)
        self.assertIn("total_limit=$((20 * 10#${PER_SHARD_LIMIT} * 10#${CHECKPOINT_BATCHES}))", workflow)
        self.assertIn("queue-${SHARD_INDEX}.json", workflow)
        self.assertIn('--queue "${queue}"', workflow)
        self.assertIn("Convert assigned queue", workflow)
        self.assertIn("Reader metadata request failed; retrying", workflow)
        self.assertIn("READER_CHM_COMMAND_TIMEOUT:", workflow)
        self.assertIn("READER_DJVU_COMMAND_TIMEOUT:", workflow)
        self.assertIn("Source metadata request failed; retrying", workflow)
        self.assertIn("for attempt in 1 2 3 4 5", workflow)
        self.assertIn("delay=$((30 * 2 ** (attempt - 1)))", workflow)
        self.assertEqual(workflow.count("retry_reader_command python scripts/scan_reader_assets.py"), 1)
        self.assertIn("prepare_reader_assets_plan.py", workflow)
        self.assertIn("stale_count", workflow)
        self.assertIn("if: inputs.dry_run != true", workflow)
        self.assertIn("max-parallel: 20", workflow)
        self.assertIn("shard: ${{ fromJSON(needs.plan.outputs.shards) }}", workflow)
        self.assertIn("fail-fast: false", workflow)
        self.assertIn("reader-assets-plan", workflow)
        self.assertIn("name: reader-assets-bundles-${{ matrix.shard }}", workflow)
        self.assertIn('bundle="output/reader-assets/bundle-${SHARD_INDEX}-1"', workflow)
        self.assertIn("path: ${{ runner.temp }}/shard-${{ matrix.shard }}.tgz", workflow)
        self.assertIn("Package conversion bundles", workflow)
        self.assertIn("SHARD_INDEX: ${{ matrix.shard }}", workflow)
        self.assertIn("name: reader-assets-bundles-${{ matrix.shard }}", workflow)
        self.assertIn("for archive in output/reader-assets/bundles/*.tgz output/reader-assets/bundles/*/*.tgz; do", workflow)
        self.assertIn('archive="${RUNNER_TEMP}/shard-${SHARD_INDEX}.tgz"', workflow)
        self.assertIn('-C "${GITHUB_WORKSPACE}/output/reader-assets"', workflow)
        self.assertIn('test -s "${archive}"', workflow)
        self.assertIn('stat "${archive}"', workflow)
        self.assertIn("Publish shard ${{ matrix.shard }}", workflow)
        self.assertIn("shopt -s nullglob", workflow)
        self.assertIn("for archive in output/reader-assets/bundles/*.tgz", workflow)
        self.assertIn("needs: [plan, convert, publish]", workflow)
        self.assertIn("needs.convert.result == 'success'", workflow)
        self.assertIn("needs.convert.result == 'failure'", workflow)
        self.assertIn("needs.convert.result == 'skipped'", workflow)
        self.assertIn('exit "${conversion_status}"', workflow)
        self.assertIn('--repo "${INPUT_REPO}"', workflow)
        self.assertNotIn('--repo "${{ inputs.repo', workflow)
        self.assertIn("conversion_status=0", workflow)
        convert_section, publish_section = workflow.split("\n  publish:\n", 1)
        self.assertNotIn("publish_reader_assets.py", convert_section.split("\n  convert:\n", 1)[1])
        self.assertIn('python scripts/publish_reader_assets.py --bundles "${bundles[@]}"', publish_section)
        self.assertIn("timeout-minutes: 360", publish_section)

if __name__ == "__main__":
    unittest.main()
